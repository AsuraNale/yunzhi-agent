# -*- coding: utf-8 -*-
"""第四轮(第三轮复核 + Codex 重跑)的修复:每条一个反例 + 改对后的对照。

高:spawn_agent 不写 fork_turns 默认 "all" · 能自己改的错误却停下来让用户打「继续」· 没引用的失效链接卡死交付前检查
中:answer 返回前就说了结果 · 资料卡片与原文的核对只查字段在不在 · Word 版丢了标题后紧贴的那一段 ·
    从没读到的来源写「查阅于」· 原生卡上选了一项又写了意见
低:任务计划再确认一次后同一个缺口又问一遍 · empty / gap 两个名字 · 记过打不开的网址卡上改写成 ok ·
    带页码、表号的行号被拒 · Word 的时间和标题级别 · PROGRESS.md · 「右侧已打开」与文字卡里的地址
"""
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import datetime, timedelta, timezone

import pytest

import card
import fetches
import liveness
import progress
import review
import yzlib
from conftest import MAIN, ROOT, jsonl, read, record_decision_answer, write
from test_round3 import (begin_round, brief_nonce, card_path, edit_card, edit_doc, good_inputs, prepare_dossier,
                         put_inputs, skill, agents, drop_decisions)
from yzlib import pl, UsageError

PY = [sys.executable, "-X", "utf8", "-B"]


def ans(card_id, *strings):
    return json.dumps({"answers": {card_id: {"answers": list(strings)}}}, ensure_ascii=False)


def cite_items(project):
    """跑 cite_check.py → 各项结果 [{check, status, detail}]。"""
    out = os.path.join(project, "records", "_cite.json")
    yzlib.run_toolkit("cite_check.py", ["--project", project, "--json", out])
    with io.open(out, encoding="utf-8") as f:
        items = json.load(f)["items"]
    os.remove(out)
    return items


def failing(items, check):
    return [i for i in items if i["check"] == check and i["status"] == "FAIL"]


# ================================================================ 高 1 · 复核助手:fork_turns 明写 "none",随机码只在说明里

def test_review_start_never_shows_the_nonce_and_asks_for_fork_turns_none(world):
    project = world("S5")
    out, record = begin_round(project, 3)
    nonce = brief_nonce(project, 3)
    assert nonce and nonce not in json.dumps(out, ensure_ascii=False)              # 主助手的对话里看不到随机码
    assert "nonce" not in record and record["nonce_sha256"] == hashlib.sha256(nonce.encode("ascii")).hexdigest()
    assert out["spawn"] == {"message": out["spawn"]["message"], "fork_turns": "none"}
    assert 'fork_turns 明写 "none"' in out["next"] and '默认是 "all"' in out["next"]
    assert "默认是 `\"all\"`" in skill("cite-trace") and "默认是 `\"all\"`" in agents()


def test_review_docs_say_how_strong_the_checks_are():
    assert "防不了故意作假" in agents() and "随机码" in agents() and "会话记录" in agents()
    with io.open(os.path.join(ROOT, "agent-tools", "review.py"), encoding="utf-8") as f:
        doc = f.read()
    assert "它**不**证明读的人没带写作对话" in doc and "fork_turns" in doc


# ================================================================ 高 2 · 自己能改的错误自己改;命令在工作区根目录下跑

def test_scripts_refuse_to_run_outside_the_workspace_root(tmp_path):
    env = dict(os.environ, YUNZHI_PROJECTS=str(tmp_path), PYTHONIOENCODING="utf-8")
    parent = os.path.dirname(ROOT)
    p = subprocess.run(PY + [os.path.join(ROOT, "agent-tools", "card.py"), "mode"], cwd=parent, env=env, capture_output=True)
    out = json.loads(p.stdout.decode("utf-8"))
    assert p.returncode == 2 and not out["ok"] and "工作区根目录" in out["error"] and "不用问用户" in out["error"]
    p = subprocess.run(PY + [os.path.join(ROOT, "agent-tools", "env_check.py")], cwd=parent, env=env, capture_output=True)
    out = json.loads(p.stdout.decode("utf-8"))
    assert p.returncode == 1 and out["ok"] is False and "工作区根目录" in out["next"] and "say" not in out
    p = subprocess.run(PY + [os.path.join(ROOT, "agent-tools", "card.py"), "mode"], cwd=ROOT, env=env, capture_output=True)
    assert p.returncode == 0 and json.loads(p.stdout.decode("utf-8"))["mode"] in ("native", "text")


def test_crash_is_exit_3_and_refusal_is_fix_it_yourself():
    def boom(_argv):
        raise ZeroDivisionError("x")
    assert yzlib.run_main(boom, ["x"]) == 3
    def usage(_argv):
        raise UsageError("参数不对")
    assert yzlib.run_main(usage, ["x"]) == 2


def test_agents_stop_rule_fixes_and_retries_without_asking():
    text = agents()
    assert "同一步最多自己重试 3 次" in text and "不让用户打「继续」" in text and "退出码 3" in text
    assert "脚本出错、或检查结果看起来不对时，不改代码：停下来" not in text      # 原来那条和「照 refuse 改」打架
    assert text.index("每条命令都在工作区根目录下跑") < text.index("## 〇、")      # 放在最前面


# ================================================================ 高 3 · 只有成稿里引用了的失效链接才要 ⚠

def add_reference(project, entry):
    path = os.path.join(project, "references.yaml")
    data = pl.load_references(path)
    data.append(entry)
    import yaml
    write(path, yaml.safe_dump({"kind": "references", "entries": data}, allow_unicode=True, sort_keys=False, width=200))


def test_uncited_dead_link_does_not_block_and_cited_one_needs_a_mark(world):
    project = world("S5")
    assert not failing(cite_items(project), "liveness")
    # 只记在资料缺口卡上、成稿里没引用的网址,后来 404:原来 ⑦ 要它带 ⚠,永远过不了
    add_reference(project, {"id": "R900", "title": "某县统计公报（打不开）", "url": "https://example.org/gone-404", "tier": "A",
                            "fetch_state": "blocked", "card": "S17", "cards": ["S17"],
                            "liveness": {"state": "failed", "detail": "4xx", "http_status": 404, "probed_at": "2026-10-05T15:00:00-04:00",
                                         "probed_from": "CA-ON"}})
    assert not failing(cite_items(project), "liveness")
    # 反例:成稿里引用了的那条变成 404,成稿没带 ⚠ → 照样拦
    refs = os.path.join(project, "references.yaml")
    cited = pl.number_entries(pl.read_md(os.path.join(project, "drafts", "%s.md" % MAIN))[1], pl.load_references(refs))[1][0][1]
    import yaml
    data = pl.load_references(refs)
    for e in data:
        if e["id"] == cited["id"]:
            e["liveness"] = {"state": "failed", "detail": "4xx", "http_status": 404, "probed_at": "2026-10-05T15:00:00-04:00",
                             "probed_from": "CA-ON"}
    write(refs, yaml.safe_dump({"kind": "references", "entries": data}, allow_unicode=True, sort_keys=False, width=200))
    assert failing(cite_items(project), "liveness")


# ================================================================ 中 4 · answer 返回前不说结果

def test_card_skill_and_agents_keep_the_first_message_neutral(world):
    assert "在 `answer` 跑完之前，不对用户说结果" in skill("yunzhi-card") and "「收到，我记一下。」" in skill("yunzhi-card")
    assert "第一件事是写回答文件、跑 `card.py answer`" in agents()
    world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    assert "answer 返回之前不对用户说结果" in out["next"]
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], "确认，开始收集资料"))
    assert res["say"].startswith("好，任务计划按第 2 版定下了")                  # 结果的说法由脚本给


def test_not_yet_has_a_say_too(world):
    world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], "先不确认"))
    assert res["result"] == "not_approved" and "你想改哪里" in res["say"]


# ================================================================ 中 5 · 卡片与原文的核对要核内容

def seal_with(project, n, inputs, findings=None):
    rd = os.path.join(project, "review")
    if findings is not None:
        write(os.path.join(rd, "findings-%d.json" % n), json.dumps(findings, ensure_ascii=False))
    put_inputs(project, n, inputs)
    return review.seal(project, n)


def test_seal_needs_a_must_fix_for_every_mismatch_then_seals(world):
    project = world("S5")
    _out, record = begin_round(project, 3)
    good = good_inputs(project, 3, record)
    with io.open(os.path.join(project, "review", "findings-3.json"), encoding="utf-8") as f:
        findings = json.load(f)
    bad_row = dict(good["source_checks"][0], result="不符", note="原文是全省数，卡上写成了分县数")
    rows = [bad_row] + good["source_checks"][1:]
    res, rc = seal_with(project, 3, dict(good, source_checks=rows))
    assert rc == 1 and "写了不符" in res["refuse"] and "没写 finding" in res["refuse"]
    tone = next(i for i, f in enumerate(findings, start=1) if f["severity"] != "must_fix")
    res, rc = seal_with(project, 3, dict(good, source_checks=[dict(bad_row, finding=tone)] + rows[1:]))
    assert rc == 1 and "不是必须改" in res["refuse"]
    must = {"severity": "must_fix", "round": 3, "status": "open", "location": "第二节第一段",
            "original": "2024 年，H 省县域公共充电桩达到 2.3 万台", "supported": "原文是全省数，不分县",
            "fix": "改成全省县域合计，并写明口径"}
    findings2 = findings + [must]
    res, rc = seal_with(project, 3, dict(good, source_checks=[dict(bad_row, finding=len(findings2))] + rows[1:]), findings2)
    assert rc == 0 and res["sealed"], res
    with io.open(os.path.join(project, "review", "result.json"), encoding="utf-8") as f:
        assert json.load(f)["counts"]["must_fix"] == 1                         # 不符进了结果:交付卡出不来


def test_seal_refuses_all_unverifiable(world):
    project = world("S5")
    _out, record = begin_round(project, 3)
    good = good_inputs(project, 3, record)
    rows = [dict(r, result="无法验证", note="原文打不开") for r in good["source_checks"]]
    res, rc = seal_with(project, 3, dict(good, source_checks=rows))
    assert rc == 1 and "全写了无法验证" in res["refuse"]
    rows[0] = good["source_checks"][0]
    res, rc = seal_with(project, 3, dict(good, source_checks=rows))
    assert rc == 0 and res["sealed"], res


def test_direction_must_agree_with_the_stance(world):
    project = world("S3")
    keep = read(card_path(project, "S10"))
    edit_card(project, "S10", lambda m: m.__setitem__("direction", "支持：农村电网改造覆盖乡镇数翻番"))
    out, rc = prepare_dossier()
    assert rc == 1 and "S10" in out["refuse"] and "两样要一致" in out["refuse"]
    edit_card(project, "S10", lambda m: m.__setitem__("direction", "反驳：农村电网改造覆盖乡镇数翻番"))
    out, rc = prepare_dossier()
    assert rc == 1 and "开头要写" in out["refuse"]
    write(card_path(project, "S10"), keep)
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_support_and_counter_need_a_concrete_data_point(world):
    project = world("S3")
    keep = read(card_path(project, "S10"))
    edit_card(project, "S10", lambda m: [f.__setitem__("value", "电网改造覆盖了大部分乡镇") for f in m["facts"]])
    out, rc = prepare_dossier()
    assert rc == 1 and "S10" in out["refuse"] and "没有一个具体的数" in out["refuse"]
    edit_card(project, "S10", lambda m: [f.__setitem__("value", "电网改造覆盖了三成乡镇") for f in m["facts"]])   # 汉字数目带单位也算
    out, rc = prepare_dossier()
    assert rc == 0, out
    write(card_path(project, "S10"), keep)


def test_method_notes_cannot_be_counter_evidence(world):
    project = world("S3")
    keep = read(card_path(project, "S10"))
    edit_card(project, "S10", lambda m: m.__setitem__("direction", "不支持：统计口径只算公共桩，不含私人桩"))
    out, rc = prepare_dossier()
    assert rc == 1 and "方法说明" in out["refuse"] and "标成背景资料" in out["refuse"]
    write(card_path(project, "S10"), keep)
    out, rc = prepare_dossier()
    assert rc == 0, out


# ================================================================ 中 6 · Word 版不丢标题后紧贴的段落;各节引用数两种格式各自核

def glue_headings(draft):
    text = read(draft)
    glued = re.sub(r"(^##+ [^\n]+)\n\n", r"\1\n", text, flags=re.M)
    assert glued != text
    write(draft, glued)


def render_both(project):
    draft = os.path.join(project, "drafts", "%s.md" % MAIN)
    refs = os.path.join(project, "references.yaml")
    for script, ext, extra in (("render_html.py", ".html", ["--cards", os.path.join(project, "cards")]), ("render_docx.py", ".docx", [])):
        rc, so, se = yzlib.run_toolkit(script, [draft, refs, os.path.join(project, "out", MAIN + ext)] + extra)
        assert rc == 0, so + se


def test_word_keeps_the_paragraph_right_after_a_heading(world):
    import cite_check
    project = world("S5")
    draft = os.path.join(project, "drafts", "%s.md" % MAIN)
    glue_headings(draft)
    render_both(project)
    entries = pl.load_references(os.path.join(project, "references.yaml"))
    want, _premise = cite_check.draft_cites_by_section(pl.read_md(draft)[1], entries)
    got = cite_check.docx_cites_by_section(os.path.join(project, "out", MAIN + ".docx"))
    assert [n for _s, n in got] == [n for _s, n in want] and sum(n for _s, n in want) == 42
    assert not failing(cite_items(project), "跨格式相等")


def test_cross_format_check_counts_citations_per_section(world):
    project = world("S5")
    draft = os.path.join(project, "drafts", "%s.md" % MAIN)
    full = read(draft)
    # 反例:Word 版少了一段(模拟原来那个丢段的生成脚本)。挑一段:它引用的每一条别处也引用了 ——
    # 去掉它,两种格式的外链集合照样相等(原来的 ④ 照样过),只有各节的引用数对不上
    paras = full.split("\n\n")
    ids = lambda text: set(re.findall(r"\[\[(R\d+)\]\]", text))

    def removable(i, p):
        if p.lstrip().startswith("#") or not ids(p):
            return False
        rest = "\n\n".join(paras[:i] + paras[i + 1:])
        return ids(p) <= ids(rest)
    k = next(i for i, p in enumerate(paras) if removable(i, p))
    refs = os.path.join(project, "references.yaml")
    write(draft, "\n\n".join(paras[:k] + paras[k + 1:]))
    yzlib.run_toolkit("render_docx.py", [draft, refs, os.path.join(project, "out", MAIN + ".docx")])
    write(draft, full)
    items = [i for i in cite_items(project) if i["check"] == "跨格式相等"]
    assert any(i["status"] == "PASS" and "外链集合" in i["detail"] for i in items)          # 原来那一项看不出来
    assert any(i["status"] == "FAIL" and "各节文内引用数与底稿对不上" in i["detail"] and ".docx" in i["detail"] for i in items)
    render_both(project)
    assert not failing(cite_items(project), "跨格式相等")


# ================================================================ 中 7 · 从没读到的来源写「未能打开」;现在打得开就先回去读

def unique_cited_source(project):
    """一张资料卡片的一个来源:它的网址只在这一张卡上,成稿里又引用了它(参考资料里会列出来)→ (卡号, 网址)。"""
    from card_check import load_cards
    refs = pl.load_references(os.path.join(project, "references.yaml"))
    cited = {e.get("url") for _no, e in pl.number_entries(pl.read_md(os.path.join(project, "drafts", "%s.md" % MAIN))[1], refs)[1]}
    owners = {}
    for c, _fn in load_cards(os.path.join(project, "cards")):
        for s in c.get("sources") or []:
            owners.setdefault(s.get("url"), set()).add(c.get("uid"))
    for c, _fn in load_cards(os.path.join(project, "cards")):
        srcs = c.get("sources") or []
        if c.get("stance") != "gap" and srcs and srcs[0].get("url") in cited and owners[srcs[0]["url"]] == {c.get("uid")}:
            return c["uid"], srcs[0]["url"]
    raise AssertionError("夹具里找不到只在一张卡上、又被引用的来源")


def test_refs_add_carries_fetch_state_and_never_read_sources_are_marked(world):
    project = world("S5")
    uid, url = unique_cited_source(project)
    edit_card(project, uid, lambda m: m["sources"][0].__setitem__("fetch_state", "blocked"))
    refs = os.path.join(project, "references.yaml")
    rc, so, se = yzlib.run_toolkit("refs_add.py", [os.path.join(project, "cards"), refs, "--merge", refs])
    assert rc == 0, so + se
    entries = {e["url"]: e for e in pl.load_references(refs)}
    assert entries[url]["fetch_state"] == "blocked"
    assert any(e.get("fetch_state") == "ok" for u, e in entries.items() if u != url)
    render_both(project)
    html = read(os.path.join(project, "out", MAIN + ".html"))
    html_refs = html[html.index('<ol class="doc-ref-list">'):]
    import docx
    paras = [p.text for p in docx.Document(os.path.join(project, "out", MAIN + ".docx")).paragraphs]
    title = entries[url]["title"]
    at = html_refs.index(title)
    li = html_refs[at:html_refs.index("</li>", at)]
    assert "未能打开" in li and "查阅于" not in li
    line = next(p for p in paras if title in p and "·" in p)
    assert "未能打开" in line and "查阅于" not in line
    assert "查阅于" in html_refs and any("查阅于" in p for p in paras)            # 读到过的照样写查阅日期


def test_refs_add_merge_rule():
    import refs_add
    assert refs_add.merged_fetch_state("blocked", "ok") == "ok" and refs_add.merged_fetch_state("ok", "failed") == "ok"
    assert refs_add.merged_fetch_state("empty", "blocked") == "blocked" and refs_add.merged_fetch_state(None, "failed") == "failed"


def test_liveness_asks_to_retry_a_source_that_was_never_read(world, monkeypatch):
    project = world("S5")
    uid, url = unique_cited_source(project)
    edit_card(project, uid, lambda m: m["sources"][0].__setitem__("fetch_state", "failed"))
    monkeypatch.setattr(liveness, "probe", lambda u, timeout: ("ok", None, 200))
    out = liveness.run(project, "CA-ON", 1)
    assert [r["url"] for r in out["retry"]] == [url]
    assert "写进成稿之前，先用网页工具" in out["next"] and "--state ok" in out["next"]
    edit_card(project, uid, lambda m: m["sources"][0].__setitem__("fetch_state", "ok"))
    out = liveness.run(project, "CA-ON", 1)
    assert "retry" not in out


# ================================================================ 中 8 · 原生卡:选了一项又写了意见

def test_option_plus_user_note_confirms_with_the_note_as_quote(world):
    project = world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    note = "时间段改成 2021 年开始，其余可以；另外请把乡镇站点的口径在卡上写得更清楚一些，方便我抽查。"
    res, rc = card.answer(MAIN, out["card_id"], ans(out["card_id"], "确认，开始收集资料", "user_note: " + note))
    assert rc == 0 and res["result"] == "approved" and res["note"] == note and note in res["next"]
    ap = pl.load_md(os.path.join(project, "task_plan.md"))[0]["approval"]
    assert ap["status"] == "approved" and ap["approval_quote"] == note[:40] and len(note) > 40
    log = jsonl(os.path.join(project, "records", "cards.jsonl"))
    assert log[-1]["note"] == note and log[-1]["choice"] == "confirm"           # 整段意见照记
    assert note in [m["text"] for m in jsonl(os.path.join(project, "records", "messages.jsonl"))]


def test_note_first_then_label_and_not_yet_with_note(world):
    project = world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], "user_note: 时间段再看看", "先不确认"))
    assert res["result"] == "not_approved" and res["note"] == "时间段再看看" and "我按你的意见改" in res["say"]
    assert pl.load_md(os.path.join(project, "task_plan.md"))[0]["approval"]["status"] != "approved"


@pytest.mark.parametrize("strings", [
    ["确认，开始收集资料", "先不确认"],                      # 两个选项名
    ["确认，开始收集资料", "时间段再看看"],                  # 第二串不是 user_note
    ["随便写的", "user_note: 时间段再看看"],                 # 两串都不是选项名
    ["确认，开始收集资料", "user_note: a", "user_note: b"],  # 三串
])
def test_other_multi_string_shapes_stay_unclear(world, strings):
    project = world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], *strings))
    assert res["result"] == "unclear"
    assert pl.load_md(os.path.join(project, "task_plan.md"))[0]["approval"]["status"] != "approved"


def test_option_plus_note_on_a_decision_card(world):
    world("S3")
    dec, rc = card.prepare("decision", MAIN, {"question": "两个缺口怎么处理？", "why": "2021 年分县数据和乡镇桩使用率都查不到。",
                                              "options": [{"label": "先写已有的部分", "effect": "缺的写进资料缺口。", "changes_plan": False},
                                                          {"label": "改研究范围", "effect": "去掉需要这两项的判断。", "changes_plan": True}]})
    assert rc == 0, dec
    res, _ = card.answer(MAIN, dec["card_id"], ans(dec["card_id"], "先写已有的部分", "user_note: 缺口写清楚查过哪些地方"))
    assert res["result"] == "selected" and res["selected"] == 1 and "缺口写清楚查过哪些地方" in res["next"]


# ================================================================ 低 · 任务计划再确认一次,同一批缺口的决定照样算

def reapprove_plan_later(project, at):
    """模拟任务计划后来又确认了一次:旧的确认记录挪进 revision_log(stamp.py --invalidate 那样),新的确认时刻更晚。
    approval、revision_log 都是过程字段,内容 hash 不变 —— 确认照样有效。"""
    def fn(m):
        old = dict(m["approval"])
        m["revision_log"] = (m.get("revision_log") or []) + [{"v": m["version"], "at": at, "who": "agent", "what": "印章作废(准备修改)",
                                                              "why": "测试", "prev_hash": old.get("approved_hash"), "prev_approval": old}]
        m["approval"]["approved_at"] = at
    edit_doc(project, "task_plan.md", fn)


def test_gap_decision_survives_a_later_plan_reapproval_if_it_covers_the_same_gaps(world):
    project = world("S3")
    drop_decisions(project)
    log = os.path.join(project, "records", "cards.jsonl")
    decided_at = "2026-10-01T17:05:00-04:00"
    record_decision_answer(project, "decision-gaps", "先写已有的部分", decided_at)
    reapprove_plan_later(project, "2026-10-02T15:00:00-04:00")
    out, rc = prepare_dossier()
    assert rc == 1 and "先用决定卡问用户" in out["refuse"]          # 没记下管哪些缺口的老决定:不算

    def set_covered(uids):
        events = jsonl(log)
        decision = next(e for e in events if e.get("type") == "answered" and e.get("card_id") == "decision-gaps")
        decision["gap_cards"] = uids                                  # 新的回答都记 gap_cards(answer 里写的)
        write(log, "".join(json.dumps(e, ensure_ascii=False) + "\n" for e in events))
    set_covered(card.gap_card_uids(project))
    out, rc = prepare_dossier()
    assert rc == 0, out
    set_covered(card.gap_card_uids(project)[:1])                      # 后来多出一个缺口:那次决定没管到它
    out, rc = prepare_dossier()
    assert rc == 1 and "先用决定卡问用户" in out["refuse"]


def test_answer_records_the_gap_cards_a_decision_covers(world):
    project = world("S3")
    dec, _ = card.prepare("decision", MAIN, {"question": "两个缺口怎么处理？", "why": "2021 年分县数据和乡镇桩使用率都查不到。",
                                             "options": [{"label": "先写已有的部分", "effect": "缺的写进资料缺口。", "changes_plan": False},
                                                         {"label": "改研究范围", "effect": "去掉需要这两项的判断。", "changes_plan": True}]})
    card.answer(MAIN, dec["card_id"], ans(dec["card_id"], "先写已有的部分"))
    assert jsonl(os.path.join(project, "records", "cards.jsonl"))[-1]["gap_cards"] == ["S17", "S18"]


# ================================================================ 低 · 一个名字:网页打开了没有数据 = empty;记过打不开的不许改写成 ok

def test_a_page_without_the_data_is_called_empty_everywhere(world):
    project = world("S3")
    src = {"url": "https://example.org/county-yearbook", "title": "统计年鉴", "tier": "A", "locator": None,
           "as_of": "2026-10-01", "fetch_state": "empty"}
    edit_card(project, "S17", lambda m: m.__setitem__("sources", [src]))
    dossier = os.path.join(project, "dossier.md")
    out, rc = prepare_dossier()
    assert rc == 1 and "empty（查了但没有内容）" in out["refuse"]          # 资料汇编还写着 gap:对不上,改法写明了名字
    write(dossier, read(dossier).replace("state: gap, owner: null}\n  - {id: G2", "state: empty, owner: null}\n  - {id: G2", 1))
    out, rc = prepare_dossier()
    assert rc == 0, out
    text = skill("evidence-card")
    assert "→ `empty`（查了但没有内容）—— 和 `fetches.py` 记的同一个名字" in text and "--state empty" in text
    assert "打开了但没有这项数据 → 资料汇编里这个缺口写 `gap`" not in text            # 原来那句把同一种情况叫成 gap


def test_logged_blocked_url_cannot_be_relabeled_ok_without_a_retry(world):
    project = world("S3")
    url = "https://example.org/blocked-then-read"
    fetches.add(project, url, "blocked", "分县数据")
    edit_card(project, "S01", lambda m: m["sources"].append({"url": url, "title": "补充公报", "tier": "A", "locator": "第 2 页",
                                                              "excerpt": "（虚构原文）补充数据。", "as_of": "2026-10-03",
                                                              "fetch_state": "ok"}))
    out, rc = prepare_dossier()
    assert rc == 1 and "卡上却写成了 ok" in out["refuse"] and "--state ok" in out["refuse"]
    fetches.add(project, url, "ok")
    out, rc = prepare_dossier()
    assert rc == 0, out
    assert "没取到的资料" not in out["message"]                              # 后来取到了,不再算「没取到」


def test_fetches_ok_only_after_a_failure(world, capsys):
    project = world("S3")
    with pytest.raises(UsageError):
        fetches._main(["fetches.py", "add", MAIN, "--url", "https://example.org/never-logged", "--state", "ok"])
    fetches.add(project, "https://example.org/never-logged", "failed")
    assert fetches._main(["fetches.py", "add", MAIN, "--url", "https://example.org/never-logged", "--state", "ok"]) == 0
    assert fetches.latest_states(project)["https://example.org/never-logged"] == "ok"


# ================================================================ 低 · 定位:带页码、表号的行号照收

@pytest.mark.parametrize("loc, bad", [("第 37 页第 5 行", False), ("表3第2行", False), ("第3行政区", False), ("p.12 line 5", False),
                                      ("Table 2, line 4", False), ("第127至138行（页面抓取）", True), ("第 5 行的数据", True),
                                      ("第 12 行", True), ("lines 12-30", True)])
def test_locator_rule_accepts_real_positions(loc, bad):
    assert bool(card.locator_problem(loc)) is bad


# ================================================================ 低 · Word:时间写 UTC,标题从标题 1 起

def test_word_time_is_utc_and_headings_start_at_level_1(world):
    import docx
    project = world("S5")
    out = os.path.join(project, "out", "检查用.docx")
    rc, so, se = yzlib.run_toolkit("render_docx.py", [os.path.join(project, "drafts", "%s.md" % MAIN),
                                                      os.path.join(project, "references.yaml"), out])
    assert rc == 0, so + se
    with zipfile.ZipFile(out) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
    stamp = re.search(r"<dcterms:created[^>]*>([^<]+)</dcterms:created>", core).group(1)
    written = datetime.strptime(stamp, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
    assert abs(datetime.now(timezone.utc) - written) < timedelta(minutes=5)
    styles = [p.style.name for p in docx.Document(out).paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert styles[0] == "Title" and styles[1] == "Heading 1" and "Heading 2" in styles
    first_h2 = styles.index("Heading 2")
    assert "Heading 1" in styles[1:first_h2]                                   # 不跳级


# ================================================================ 低 · PROGRESS.md 跟着走,确认都记一行

def test_progress_md_follows_advance_verifying_and_logs_the_dossier_approval(world):
    project = world("S3")
    out, rc = prepare_dossier()
    assert rc == 0, out
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], out["ask"]["questions"][0]["options"][0]["label"]))
    assert res["result"] == "approved"
    text = read(os.path.join(project, "PROGRESS.md"))
    assert re.search(r"^- \d{1,2}-\d{2} 资料汇编第 1 版已确认（由助手记录）$", text, re.M), text
    # 助手自己推进到交付前检查之后,刷新进度页就把第一行对上
    assert yzlib.run_toolkit("stamp.py", [os.path.join(project, "task_plan.md"), "--advance", "drafting"])[0] == 0
    assert yzlib.run_toolkit("stamp.py", [os.path.join(project, "task_plan.md"), "--advance", "verifying"])[0] == 0
    progress.write_page(project)
    first = read(os.path.join(project, "PROGRESS.md")).split("\n")[0]
    assert re.fullmatch(r"%s · verifying · \d{4}-\d{2}-\d{2}" % re.escape(MAIN), first), first


# ================================================================ 低 · 「右侧已打开」只在挂过右侧时写;文字卡不放 file:///

def test_pane_open_flag_and_text_card_location(world, monkeypatch):
    project = world("S1")
    out, rc = card.prepare("task_plan", MAIN, {"outside": []})
    assert rc == 0 and "右侧已打开" not in out["message"] and "[点开看](<file:///" in out["message"]
    out, rc = card.prepare("task_plan", MAIN, {"outside": []}, pane_open=True)
    assert rc == 0 and "（右侧已打开 · [点开看](<file:///" in out["message"]
    monkeypatch.setenv("YUNZHI_CARDS", "text")
    out, rc = card.prepare("task_plan", MAIN, {"outside": []}, pane_open=True)
    assert rc == 0 and "file:///" not in out["fallback_text"] and "file:///" not in out["message"]
    assert "（右侧已打开 · 全文在项目文件夹的「查看」文件夹里，网页「任务计划」）" in out["fallback_text"]
    assert "--pane-open" in skill("yunzhi-card") and "--pane-open" in skill("yunzhi-progress")

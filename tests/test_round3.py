# -*- coding: utf-8 -*-
"""第三轮(10-03 用户在 Codex 桌面版里完整跑了一遍,复核员逐条审)的修复:每条一个反例 + 改对后的对照。

高:复核不独立(带着写作对话)· 交付卡前那条消息没发 · 任务计划写的「停下来问」被跳过 · 反面证据漏了、读反了
中:卡号写错就认不出 · 连着两个「1」被吞 · 打不开没留痕、写成没有这项数据、缺口数对不上 · 重新确认太重 ·
    提纲每一版不上右侧 · 快照没贴 · 定位写成行号 · 复核不逐段
低:PROGRESS.md 停在 planning · 委托原话被改 · 交付前检查永远「另有 1 项提醒」· 沙盒不能联网时链接全记失败 ·
    Word 版的参考资料和文档属性 · 查看页把查阅日期标成发布于
"""
import io
import json
import os
import re

import pytest

import card
import fetches
import liveness
import progress
import projects
import review
import view
import yzlib
from conftest import MAIN, ROOT, jsonl, read, record_decision_answer, upgrade_cards, write
from yzlib import pl, UsageError

SKILLS = os.path.join(ROOT, ".agents", "skills")


def skill(name):
    return read(os.path.join(SKILLS, name, "SKILL.md"))


def agents():
    return read(os.path.join(ROOT, "AGENTS.md"))


def card_path(project, uid):
    return os.path.join(project, "cards", "%s.md" % uid)


def edit_card(project, uid, fn):
    meta, body = pl.read_md(card_path(project, uid))
    fn(meta)
    pl.write_md(card_path(project, uid), meta, body)


def edit_doc(project, doc, fn):
    path = os.path.join(project, doc)
    meta, body = pl.read_md(path)
    fn(meta)
    pl.write_md(path, meta, body)


def prepare_dossier():
    return card.prepare("dossier", MAIN, {"outside": []})


def ans(card_id, label):
    return json.dumps({"answers": {card_id: {"answers": [label]}}}, ensure_ascii=False)


# ================================================================ 高 1 · 独立复核不带写作对话

def begin_round(project, n):
    """主助手开始第 n 轮(review.py start),再照「照说明做完的复核助手」写好发现与报告(上一轮的发现照抄过来)。"""
    out = review.start(project, n)
    k = review.last_sealed_round(project, n)
    with io.open(os.path.join(project, "review", "result-%d.json" % k), encoding="utf-8") as f:
        findings = json.load(f)["findings"]
    write(os.path.join(project, "review", "findings-%d.json" % n), json.dumps(findings, ensure_ascii=False))
    write(os.path.join(project, "review", "report-%d.md" % n), "# 第 %d 轮复核报告\n\n逐段核过。\n" % n)
    with io.open(os.path.join(project, "review", "round-%d.json" % n), encoding="utf-8") as f:
        record = json.load(f)
    return out, record


def brief_nonce(project, n):
    """复核助手从这一轮的说明里读到的随机码(第四轮起开始记录只存它的 sha256)。"""
    m = re.search(r"随机码（nonce）：`([0-9a-f]+)`", read(os.path.join(project, "review", "brief-%d.md" % n)))
    return m.group(1)


def good_inputs(project, n, record):
    cards = review.live_cards(project)
    return {"round": n, "nonce": brief_nonce(project, n), "files_read": list(record["files"]), "history": "没有", "seen_codes": [],
            "coverage": [{"para": p["id"], "claims": 1, "result": "核过"} for p in record["paragraphs"]],
            "source_checks": [{"card": c, "result": "相符", "note": "原文对得上"} for c in cards[:record["source_checks_min"]]]}


def put_inputs(project, n, data):
    write(os.path.join(project, "review", "inputs-%d.json" % n), json.dumps(data, ensure_ascii=False))


def test_review_start_keeps_the_code_out_of_every_file(world):
    project = world("S5")
    out, record = begin_round(project, 3)
    code = re.search(r"YZC-[0-9A-F]{8}", out["code_notice"]).group(0)
    assert out["spawn"]["fork_turns"] == "none" and code not in out["spawn"]["message"]
    assert "brief-3.md" in out["spawn"]["message"] and "review-module" in out["spawn"]["message"]
    for root, _dirs, files in os.walk(project):
        for fn in files:
            with io.open(os.path.join(root, fn), "rb") as f:
                assert code.encode("ascii") not in f.read(), fn          # 暗号只在主助手这一次的输出里
    brief = read(os.path.join(project, "review", "brief-3.md"))
    assert brief_nonce(project, 3) and record["paragraphs"] and all(p["id"] in brief for p in record["paragraphs"])
    assert all("`%s`" % f in brief for f in record["files"]) and record["source_checks_min"] == 5


def test_seal_needs_evidence_of_a_clean_reviewer_then_seals(world):
    project = world("S5")
    out, record = begin_round(project, 3)
    code = re.search(r"YZC-[0-9A-F]{8}", out["code_notice"]).group(0)
    result_before = read(os.path.join(project, "review", "result.json"))
    good = good_inputs(project, 3, record)

    def refused(data, words):
        if data is None:
            os.remove(os.path.join(project, "review", "inputs-3.json")) if os.path.exists(os.path.join(project, "review", "inputs-3.json")) else None
        else:
            put_inputs(project, 3, data)
        res, rc = review.seal(project, 3)
        assert rc == 1 and not res["sealed"] and words in " ".join(res["problems"]), res
        assert read(os.path.join(project, "review", "result.json")) == result_before    # 拒绝时什么都不写
        assert not os.path.exists(os.path.join(project, "review", "sealed-3.json"))

    refused(None, "inputs 文件还没写")
    refused(dict(good, seen_codes=[code]), "带着写作对话")                                    # 分了对话的复核助手报出了暗号
    refused(dict(good, nonce="00000000"), "随机码对不上")
    refused(dict(good, files_read=good["files_read"][1:]), "files_read 少了")
    refused(dict(good, files_read=good["files_read"] + ["projects/%s/records/cards.jsonl" % MAIN]), "files_read 多了")
    refused(dict(good, coverage=good["coverage"][1:]), "coverage 少了这几段")
    refused(dict(good, source_checks=good["source_checks"][:2]), "至少要 5 张")
    put_inputs(project, 3, good)
    report = os.path.join(project, "review", "report-3.md")
    write(report, read(report) + "\n见过的东西：%s\n" % code)                                   # 暗号出现在交的文件里
    res, rc = review.seal(project, 3)
    assert rc == 1 and "出现了这一轮的「复核暗号」" in " ".join(res["problems"])
    write(report, "# 第 3 轮复核报告\n\n逐段核过。\n")
    res, rc = review.seal(project, 3)
    assert rc == 0 and res["sealed"], res
    assert os.path.isfile(os.path.join(project, "review", "sealed-3.json"))
    res, rc = review.check(project)
    assert rc == 0, res
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 0, out


def test_delivery_only_accepts_a_review_sealed_by_review_py(world):
    project = world("S5")
    sealed = os.path.join(project, "review", "sealed-2.json")
    keep = read(sealed)
    os.remove(sealed)
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 1 and "不是用 review.py seal 封存的" in " ".join(out["problems"])
    write(sealed, keep)
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 0, out
    inputs = os.path.join(project, "review", "inputs-2.json")
    good = read(inputs)
    write(inputs, good.replace('"seen_codes": []', '"seen_codes": ["YZC-1234ABCD"]'))      # 封存之后改了 inputs
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 1 and "封存之后改过" in " ".join(out["problems"])


def test_skills_spawn_the_reviewer_without_the_conversation():
    # 第四轮:Codex 0.159.2 的 spawn_agent 不写 fork_turns 默认是 "all" —— 要明写 "none",不能再说「不加」
    ct = skill("cite-trace")
    assert '$PY "agent-tools/review.py" start' in ct and '$PY "agent-tools/review.py" check' in ct
    assert '`fork_turns` 明写 `"none"`' in ct and '绝不写 `"all"`' in ct
    assert '$PY "agent-tools/review.py" seal' in skill("review-module")
    for name in os.listdir(SKILLS):
        text = skill(name)
        assert "--seal" not in text or "review_result.py" not in text.split("--seal")[0][-80:], name   # 不再直接用 toolkit 封存
        for line in text.split("\n"):
            if "fork_turns" in line:
                assert "none" in line and "不加" not in line, line
    for line in agents().split("\n"):
        if "fork_turns" in line:
            assert "none" in line and "不加" not in line, line
    assert "review.py" in agents() and '`fork_turns` 明写 `"none"`' in agents()


# ================================================================ 高 2 · 要紧的内容也在原生选项卡上

@pytest.mark.parametrize("stage, kind, fields, must", [
    ("S1", "task_plan", {"outside": []}, []),
    ("S3", "dossier", {"outside": []}, ["初步结论：部分成立", "资料卡片 18 张"]),
    ("S4", "outline", {}, ["合计 8000 字（目标 8000 字）", "没有资料卡片支撑的节：0 个"]),
    ("S5", "delivery", {}, ["交付前检查：", "独立复核：必须改的 0 处"]),
])
def test_native_card_carries_the_three_points(world, stage, kind, fields, must):
    project = world(stage)
    out, rc = card.prepare(kind, MAIN, dict(fields))
    assert rc == 0, out
    question = out["ask"]["questions"][0]["question"]
    if kind == "delivery":
        _m, body, _p = pl.load_md(card.delivery_commitments_path(project))
        items = pl.commitment_problems(body, whole=True)[2]
    else:
        items = pl.commitment_problems(pl.load_md(os.path.join(project, card.DOC_OF[kind]))[1])[2]
    for label, it in zip(card.POINT_LABELS, items):
        assert "%s：%s" % (label, it["content"]) in question
    for m in must:
        assert m in question, question


def test_flow_change_card_carries_the_change(world):
    world("S2-flow")
    out, rc = card.prepare("flow_change", MAIN, {"changes": "去掉拟定提纲这一步，写作按资料汇编的结构走。"})
    assert rc == 0 and "改动环节：去掉拟定提纲这一步" in out["ask"]["questions"][0]["question"]


def test_agents_and_card_skill_say_reread_after_compaction():
    assert "上下文被压缩" in agents() and "yunzhi-card" in agents()
    assert "上下文被压缩" in skill("yunzhi-card")


# ================================================================ 高 3 · 资料汇编列了缺口:先用决定卡问用户

def drop_decisions(project):
    log = os.path.join(project, "records", "cards.jsonl")
    write(log, "".join(json.dumps(e, ensure_ascii=False) + "\n" for e in jsonl(log) if e.get("kind") != "decision"))


def test_dossier_with_gaps_needs_a_decision_since_the_plan_was_confirmed(world):
    project = world("S3")
    drop_decisions(project)
    out, rc = prepare_dossier()
    assert rc == 1 and "先用决定卡问用户" in out["refuse"]
    record_decision_answer(project, "decision-old", "先写已有的部分", "2026-09-29T09:00:00-04:00")   # 任务计划确认之前的不算
    out, rc = prepare_dossier()
    assert rc == 1 and "先用决定卡问用户" in out["refuse"]
    # 照真的流程问一次:出决定卡,用户在卡上选了一项
    dec, rc = card.prepare("decision", MAIN, {"question": "两个缺口怎么处理？", "why": "2021 年分县数据和乡镇桩使用率都查不到。",
                                              "options": [{"label": "先写已有的部分", "effect": "缺的写进资料缺口。", "changes_plan": False},
                                                          {"label": "改研究范围", "effect": "去掉需要这两项的判断。", "changes_plan": True}]})
    assert rc == 0, dec
    res, _ = card.answer(MAIN, dec["card_id"], ans(dec["card_id"], "缺的再找找看"))     # 只写了意见、没选:不算
    assert res["result"] == "note"
    out, rc = prepare_dossier()
    assert rc == 1 and "先用决定卡问用户" in out["refuse"]
    dec, rc = card.prepare("decision", MAIN, {"question": "两个缺口怎么处理？", "why": "2021 年分县数据和乡镇桩使用率都查不到。",
                                              "options": [{"label": "先写已有的部分", "effect": "缺的写进资料缺口。", "changes_plan": False},
                                                          {"label": "改研究范围", "effect": "去掉需要这两项的判断。", "changes_plan": True}]})
    res, _ = card.answer(MAIN, dec["card_id"], ans(dec["card_id"], "先写已有的部分"))
    assert res["result"] == "selected"
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_review_zeroth_layer_checks_stop_conditions():
    text = skill("review-module")
    assert "停下来问用户的条件" in text and "你决定" in text


# ================================================================ 高 4 · 原文摘录、方向、方法说明、找反面证据

def test_cards_need_excerpts_directions_and_cited_facts(world):
    project = world("S3")
    keep = read(card_path(project, "S10"))
    edit_card(project, "S10", lambda m: m["sources"][0].pop("excerpt"))
    out, rc = prepare_dossier()
    assert rc == 1 and "S10" in out["refuse"] and "原文摘录" in out["refuse"]
    write(card_path(project, "S10"), keep)
    edit_card(project, "S10", lambda m: m.pop("direction"))
    out, rc = prepare_dossier()
    assert rc == 1 and "S10" in out["refuse"] and "和核心判断的关系" in out["refuse"]
    write(card_path(project, "S10"), keep)
    edit_card(project, "S10", lambda m: [f.__setitem__("cited", False) for f in m["facts"]])   # 方法说明标成了不支持
    out, rc = prepare_dossier()
    assert rc == 1 and "S10" in out["refuse"] and "标成背景资料" in out["refuse"]
    write(card_path(project, "S10"), keep)
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_argument_dossier_needs_a_counter_evidence_search(world):
    project = world("S3")
    path = os.path.join(project, "dossier.md")
    keep = read(path)
    edit_doc(project, "dossier.md", lambda m: m.pop("counter_search"))
    out, rc = prepare_dossier()
    assert rc == 1 and "找反面证据" in out["refuse"]
    write(path, keep)
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_evidence_skill_asks_for_excerpts_directions_and_a_counter_search():
    text = skill("evidence-card")
    assert "excerpt:" in text and "direction:" in text and "counter_search:" in text
    assert "专门找一遍反面证据" in text and "方法说明、口径说明、背景介绍不是正反证据" in text


# ================================================================ 中 5 · 卡号只差 - / _ / 大小写也认

@pytest.mark.parametrize("raw", [
    '{"answers":{"task_plan_v2":{"answers":["确认，开始收集资料"]}}}',
    '{"answers":{"Task-Plan-V2":{"answers":["确认，开始收集资料"]}}}',
    '{"answers":{"q1":{"answers":["确认，开始收集资料"]}}}',
])
def test_answer_accepts_near_ids_and_a_lone_exact_label(world, raw):
    project = world("S1")
    card.prepare("task_plan", MAIN, {"outside": []})
    res, rc = card.answer(MAIN, "task-plan-v2", raw)
    assert rc == 0 and res["result"] == "approved", res
    assert pl.load_md(os.path.join(project, "task_plan.md"))[0]["approval"]["status"] == "approved"


def test_near_id_also_carries_a_typed_note(world):
    # 上面三条的回答都正好是选项名:卡号认不出时,「只有一道题、回答正好是选项名」那一条也兜得住,
    # 所以盯不住「卡号只差 - / _ 也认」这一条(变异 Q18 改坏它,三条照样绿)。写的是意见、不是选项名时,
    # 只有卡号这一条认得出:认出 → 先不确认 + 意见;认不出 → unclear。
    project = world("S1")
    card.prepare("task_plan", MAIN, {"outside": []})
    raw = '{"answers":{"task_plan_v2":{"answers":["时间段改成 2021 年开始"]}}}'
    res, rc = card.answer(MAIN, "task-plan-v2", raw)
    assert rc == 0 and res["result"] == "not_approved" and res["note"] == "时间段改成 2021 年开始", res
    assert pl.load_md(os.path.join(project, "task_plan.md"))[0]["approval"]["status"] != "approved"


def test_card_skill_says_no_outcome_before_answer():
    assert "`answer` 跑完之前，不对用户说结果" in skill("yunzhi-card") and "不对用户说结果" in agents()


# ================================================================ 中 7 · 没取到的留痕、打不开不写成没有、缺口对得上

def test_gap_cards_and_dossier_gaps_must_agree(world):
    project = world("S3")
    keep = read(card_path(project, "S18"))
    edit_card(project, "S18", lambda m: m.__setitem__("deprecated", True))
    out, rc = prepare_dossier()
    assert rc == 1 and "对不上" in out["refuse"]
    write(card_path(project, "S18"), keep)
    dossier = os.path.join(project, "dossier.md")
    keep_d = read(dossier)
    write(dossier, keep_d.replace("state: gap, owner: null}\n  - {id: G2", "state: blocked, owner: null}\n  - {id: G2", 1))
    out, rc = prepare_dossier()
    assert rc == 1 and "打不开" in out["refuse"]                   # 资料汇编说打不开,资料缺口卡上没有打不开的记录
    edit_card(project, "S17", lambda m: m.__setitem__("sources", [{"url": "https://example.org/blocked", "title": "统计年鉴", "tier": "A",
                                                                    "locator": None, "as_of": "2026-10-01", "fetch_state": "blocked"}]))
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_logged_failed_fetches_must_leave_a_trace_on_a_card(world):
    project = world("S3")
    fetches.add(project, "https://example.org/gone", "blocked", "2021 年分县数据")
    out, rc = prepare_dossier()
    assert rc == 1 and "https://example.org/gone" in out["refuse"]
    edit_card(project, "S17", lambda m: m.__setitem__("sources", [{"url": "https://example.org/gone", "title": "统计年鉴", "tier": "A",
                                                                    "locator": None, "as_of": "2026-10-01", "fetch_state": "blocked"}]))
    dossier = os.path.join(project, "dossier.md")
    write(dossier, read(dossier).replace("state: gap, owner: null}\n  - {id: G2", "state: blocked, owner: null}\n  - {id: G2", 1))
    fetches.add(project, "https://example.org/gone", "blocked", "2021 年分县数据")     # 同一个网址又记一次:还是一个来源(第八轮)
    out, rc = prepare_dossier()
    assert rc == 0, out
    said = "没取到的资料：打不开 1 个来源"
    assert said in out["message"] and said in out["ask"]["questions"][0]["question"]
    page = read(view.render(project, "dossier")[0])          # 资料汇编的阅读页照记录说「打不开」,不说成没有这项数据
    assert "2021 年分县充电桩数（打不开）" in page and "乡镇桩使用率（没有这项数据）" in page


# ================================================================ 中 8 · 重新确认轻一点

def reopen_dossier(project):
    """S5 → 资料汇编要重新确认:作废、升版、补齐第三轮的规矩、推进到等确认;提纲不动(确认记录照旧有效)。"""
    path = os.path.join(project, "dossier.md")
    assert yzlib.run_toolkit("stamp.py", [path, "--invalidate", "--why", "补一张资料卡片"])[0] == 0

    def bump(m):
        m["version"] = max(e["v"] for e in m["revision_log"])
        m["counter_search"] = [{"query": "没参加活动的县 增速", "result": "找到 2 份，已做成不支持的卡", "found": ["S10", "S11"]}]
    edit_doc(project, "dossier.md", bump)
    upgrade_cards(project, with_direction=True)
    record_decision_answer(project, "decision-gap", "先写已有的部分", yzlib.now_iso())
    assert yzlib.run_toolkit("stamp.py", [os.path.join(project, "task_plan.md"), "--advance", "gate2_awaiting"])[0] == 0


def test_dossier_reconfirm_keeps_a_valid_outline_and_asks_to_continue(world):
    project = world("S5")
    reopen_dossier(project)
    assert yzlib.approved_valid(project, "outline.md")
    out, rc = prepare_dossier()
    assert rc == 0, out
    assert out["ask"]["questions"][0]["question"].split("\n")[0] == "资料汇编第 2 版可以定下来，继续写吗？"
    assert "重新确认 · 资料汇编" in out["message"]
    res, _ = card.answer(MAIN, out["card_id"], ans(out["card_id"], "确认，继续写"))
    assert res["result"] == "approved" and res["state"] == "drafting" and "不用再确认" in res["next"]
    assert yzlib.approved_valid(project, "outline.md")             # 提纲的确认记录照旧有效,没被要求重新确认


def test_outline_reconfirm_wording_fits_the_stage(world):
    project = world("S4")
    edit_doc(project, "outline.md", lambda m: m.__setitem__("revision_log", (m.get("revision_log") or []) + [
        {"v": m["version"], "at": "2026-10-03T15:00:00-04:00", "who": "agent", "what": "印章作废(准备修改)", "why": "测试",
         "prev_approval": {"status": "approved", "approved_hash": "000000000000"}}]))
    out, rc = card.prepare("outline", MAIN, {})
    assert rc == 0, out
    assert out["ask"]["questions"][0]["question"].split("\n")[0] == "提纲第 3 版可以定下来，继续写吗？"
    assert [o["label"] for o in out["ask"]["questions"][0]["options"]] == ["确认，继续写", "先不确认"]


def test_stage_durations_count_from_the_first_entry(world):
    project = world("S4")
    first, later = "2026-09-30T10:43:00-04:00", "2026-10-01T09:00:00-04:00"
    edit_doc(project, "task_plan.md", lambda m: m["approval"].__setitem__("approved_at", later))   # 任务计划后来又确认过一次
    log = os.path.join(project, "records", "cards.jsonl")
    with io.open(log, "a", encoding="utf-8", newline="\n") as f:
        for at in (first, later):
            f.write(json.dumps({"type": "answered", "at": at, "card_id": "task-plan", "kind": "task_plan", "result": "approved",
                                "choice": "confirm", "label": "确认，开始收集资料"}, ensure_ascii=False) + "\n")
    m = progress.model(project)
    sources = next(e for e in m["stages"] if e["id"] == "sources")
    end = progress.parse_time(pl.load_md(os.path.join(project, "dossier.md"))[0]["approval"]["approved_at"])
    assert sources["duration"] == progress.duration(progress.parse_time(first), end)
    assert sources["duration"] != progress.duration(progress.parse_time(later), end)


# ================================================================ 中 9 · 提纲每一版写完就上右侧

def test_view_outline_check_shows_each_version_on_the_pane(world, capsys):
    project = world("S4")
    assert view.main(["view.py", MAIN, "outline", "--pane", "--check"]) == 0
    out = json.loads(capsys.readouterr().out)
    from outline_check import parse_outline_md
    _meta, nodes = parse_outline_md(os.path.join(project, "outline.md"))
    assert out["check_ok"] is True and out["words_sum"] == pl.outline_total_words(nodes) and out["total_words"] == 8000
    assert yzlib.pane_kind(project) == "view:outline"
    path = os.path.join(project, "outline.md")
    write(path, read(path).replace("- evidence: [S02, S10]", "- evidence: [S02, S99]", 1))
    view.main(["view.py", MAIN, "outline", "--pane", "--check"])
    out = json.loads(capsys.readouterr().out)
    assert out["check_ok"] is False and out["check_problems"] and "先改提纲" in out["next"]


def test_outline_skill_shows_every_version():
    assert '$PY "agent-tools/view.py" "<项目名>" outline --pane --check' in skill("outline-cocreate")


# ================================================================ 中 10 · 每一轮最后贴快照 → 第五轮 M1:先停用

def test_turn_end_snapshot_is_paused():
    # 10-04:桌面版 26.930 上 visualize{…} 显示成了原始文字;说明里不再要求贴快照,写清停用的原因
    assert "每一轮的最后一步" not in agents() and "每一轮的最后一步" not in skill("yunzhi-progress")
    assert "进度快照这一版先停用" in agents() and "这一版停用" in skill("yunzhi-progress")
    assert "原始文字" in agents() and "原始文字" in skill("yunzhi-progress")


# ================================================================ 中 11 · 定位不写抓取工具的行号

@pytest.mark.parametrize("loc, bad", [("第127至138行", True), ("第 12 行", True), ("lines 12-30", True), ("L12-L30", True),
                                      ("第 4 章表 4-3", False), ("第 37 页", False), ("p.37 合并利润表", False), ("二、主要结论", False)])
def test_locator_rule(loc, bad):
    assert bool(card.locator_problem(loc)) is bad


def test_line_number_locators_are_refused_on_cards_and_references(world):
    project = world("S3")
    keep = read(card_path(project, "S01"))
    edit_card(project, "S01", lambda m: m["sources"][0].__setitem__("locator", "第127至138行"))
    out, rc = prepare_dossier()
    assert rc == 1 and "第127至138行" in out["refuse"] and "抓取工具给的行号" in out["refuse"]
    write(card_path(project, "S01"), keep)
    out, rc = prepare_dossier()
    assert rc == 0, out


def test_line_number_locator_in_references_blocks_delivery(world):
    project = world("S5")
    refs = os.path.join(project, "references.yaml")
    keep = read(refs)
    first = re.search(r"locator: '([^']*)'", keep).group(1)
    write(refs, keep.replace("locator: '%s'" % first, "locator: '第127至138行'", 1))
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 1 and "第127至138行" in " ".join(out["problems"])
    write(refs, keep)
    out, rc = card.prepare("delivery", MAIN, {})
    assert rc == 0, out


# ================================================================ 低 · PROGRESS.md、委托原话、检查项计数、联网、Word 版、查阅日期

def test_progress_md_first_line_follows_the_state(world):
    project = world("S1")
    out, _ = card.prepare("task_plan", MAIN, {"outside": []})
    card.answer(MAIN, out["card_id"], ans(out["card_id"], "确认，开始收集资料"))
    lines = read(os.path.join(project, "PROGRESS.md")).split("\n")
    assert re.fullmatch(r"%s · collecting · \d{4}-\d{2}-\d{2}" % re.escape(MAIN), lines[0]), lines[0]
    assert "# 进度" in lines[1:]                                        # 原来的笔记留着
    assert read(os.path.join(project, "PROGRESS.md")).count(" · collecting · ") == 1


def test_new_project_keeps_the_request_verbatim_and_refuses_a_title_line(projects_root):
    with pytest.raises(UsageError):
        projects.new_project("短租", "# 短租市场比较\n新建项目：比较安省和美国的短租市场。", [])
    text = "新建项目：比较安省和美国的短租市场，写一份两千字左右的说明。\n"
    out = projects.new_project("短租", text, [])
    assert read(os.path.join(projects_root, out["project"], "inputs", "request.md")) == text


def test_checks_count_leaves_out_the_client_only_item():
    items = [{"check": "无悬空占位", "status": "PASS"}, {"check": "预算仪表·步数", "status": "WARN"}, {"check": "预算仪表", "status": "PASS"}]
    assert card.checks_counts(items) == {"total": 2, "pass": 2, "warn": 0, "fail": 0}
    assert card.checks_phrase(card.checks_counts(items)) == "2 / 2 通过"


def test_liveness_offline_is_unchecked_and_web_tool_pages_stay_reachable(world, monkeypatch):
    project = world("S5")
    monkeypatch.setattr(liveness, "probe", lambda url, timeout: ("failed", "net", None))
    out = liveness.run(project, "CA-ON", 1)
    entries = pl.load_references(os.path.join(project, "references.yaml"))
    web = [e for e in entries if str(e.get("url", "")).startswith("http")]
    assert out["offline"] and out["counts"].get("unchecked") == len(web) and "未能自动检查" in out["say"]
    assert all(e["liveness"].get("unchecked") and e["liveness"]["note"].startswith("未能自动检查") for e in web)
    rid = web[0]["id"]
    # 第八轮:mark 只收有读过证据的(资料卡片上这个来源取到了、写了原文摘录);示例项目的卡片还没有摘录 → 先拒
    with pytest.raises(UsageError):
        liveness.mark(project, [rid])

    def add_excerpt(meta):
        for s in meta["sources"]:
            if s.get("url") == web[0]["url"]:
                s["excerpt"] = "（虚构原文）截至 2024 年底，全省县域公共充电桩 2.3 万台。"
    edit_card(project, web[0]["card"], add_excerpt)
    liveness.mark(project, [rid])
    out = liveness.run(project, "CA-ON", 1)                           # 还是不能联网:打开过的那条不被盖掉
    entries = {e["id"]: e for e in pl.load_references(os.path.join(project, "references.yaml"))}
    assert entries[rid]["liveness"]["state"] == "ok" and entries[rid]["liveness"]["via"] == "web-tool"
    assert out["counts"].get("ok/web-tool") == 1
    # 能联网、个别连不上:照实记失败,不当成「未能自动检查」
    monkeypatch.setattr(liveness, "probe", lambda url, timeout: ("ok", None, 200) if url == web[0]["url"] else ("failed", "net", None))
    out = liveness.run(project, "CA-ON", 1)
    assert not out["offline"] and "unchecked" not in out["counts"]


def test_word_version_has_proper_properties_heading_and_reference_details(world):
    import docx
    project = world("S5")
    draft = os.path.join(project, "drafts", "%s.md" % MAIN)
    out = os.path.join(project, "out", "检查用.docx")
    rc, so, se = yzlib.run_toolkit("render_docx.py", [draft, os.path.join(project, "references.yaml"), out])
    assert rc == 0, so + se
    d = docx.Document(out)
    props = d.core_properties
    assert props.author == "云织" and props.last_modified_by == "云织" and props.created.year >= 2026
    chapter_styles = {p.style.name for p in d.paragraphs if p.style.name.startswith("Heading") and p.text != "参考资料"}
    ref_heading = next(p for p in d.paragraphs if p.text == "参考资料")
    # 第四轮:稿件最上一级的节(##)是 Word 的标题 1,不跳级;参考资料和章同一级
    assert ref_heading.style.name in chapter_styles and ref_heading.style.name == "Heading 1"
    entries = pl.load_references(os.path.join(project, "references.yaml"))
    texts = [p.text for p in d.paragraphs]
    e = entries[0]
    line = next(t for t in texts if e["title"] in t)
    assert e["url"] in line and str(e["locator"]) in line and "查阅于 %s" % e["as_of"] in line


def test_cards_page_labels_the_fetch_date_and_shows_excerpts(world):
    project = world("S3")
    path, _title = view.render(project, "cards")
    page = read(path)
    assert "发布于" not in page and "查阅于 " in page and "原文摘录：「（虚构原文）" in page and "和核心判断的关系：" in page


def test_env_check_runs_once_per_conversation():
    assert "一个对话只跑一次" in agents() and "一个对话只跑一次" in skill("yunzhi")

# -*- coding: utf-8 -*-
"""pipeline_status 的机器窗口契约测试。

跑法: python -X utf8 test_pipeline_status.py

守两件事:
  1. ⭐ **加不加 --json,人读输出(stdout)与 exit code 逐字节不变** ——
     --json 是纯增量出口,面板要什么都不许改变人读那一路的行为。
  2. --json 写出的结构里,面板需要的事实字段齐备且与人读那行同源。
"""
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pipeline_lib import content_hash, read_md, write_md  # noqa: E402

# v1.1.0:session_facts.mjs 读 `DSH_HOME` 下的会话,且注入的会话变量无效时不再
# 回落到启发式。这里的测试都在自己的临时 HOME 里造日志 —— 外面的 DSH_* 一个都
# 不许漏进子进程(在 DSH 的 shell 工具里跑测试时,harness 会注入 DSH_HOME 与
# DSH_SESSION_ID,那样测的就是开发机上的真会话,而不是夹具)。
for _k in ("DSH_HOME", "DSH_SESSION_ID", "DSH_SESSION_JSONL"):
    os.environ.pop(_k, None)

# 换行常量:这份文件里生成内容一律用它拼,不写转义序列 ——
# 反斜杠在通往这里的那条链上被吃掉过多次。
NL = chr(10)

SCRIPT = os.path.join(HERE, "pipeline_status.py")


def _doc(path, title, approved=True):
    """写一份印章有效的正本(hash 按规格现算,免得测试自己造出 stale)。"""
    body = "\n# %s\n\n正文若干。\n" % title
    meta = {"title": title, "version": 1, "updated_at": "2020-01-01T00:00:00+08:00"}
    meta["approval"] = {"status": "approved" if approved else "draft"}
    if approved:
        meta["approval"]["approved_at"] = "2020-01-01T00:00:00+08:00"
        meta["approval"]["approved_hash"] = content_hash(meta, body)
    write_md(path, meta, body)


def _project(tmp):
    _doc(os.path.join(tmp, "task_plan.md"), "任务计划")
    _doc(os.path.join(tmp, "dossier.md"), "证据汇编")
    _doc(os.path.join(tmp, "outline.md"), "大纲")
    meta, _ = __import__("pipeline_lib").read_md(os.path.join(tmp, "task_plan.md"))
    meta["pipeline_status"] = "gate3_approved"
    meta["last_turn_at"] = "2020-01-01T00:00:00+08:00"
    body = "\n# 任务计划\n\n正文若干。\n"
    meta["approval"]["approved_hash"] = content_hash(meta, body)
    write_md(os.path.join(tmp, "task_plan.md"), meta, body)
    with io.open(os.path.join(tmp, "PROGRESS.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("对接窗口3 已过,正在写正本\n\n细节若干\n")
    return tmp


def _run(args):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    r = subprocess.run([sys.executable, "-X", "utf8", SCRIPT] + args,
                       capture_output=True, env=env)
    return r.returncode, r.stdout, r.stderr



def _run_stamp(args):
    """跑 stamp.py → (exit, stdout, stderr)。"""
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    r = subprocess.run([sys.executable, "-X", "utf8",
                        os.path.join(HERE, "stamp.py")] + args,
                       capture_output=True, env=env)
    return r.returncode, r.stdout, r.stderr

class JsonIsPurelyAdditive(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-")
        _project(self.tmp)
        self.out = os.path.join(self.tmp, "status.json")

    def test_stdout_and_exit_code_are_byte_identical(self):
        plain_code, plain_out, plain_err = _run([self.tmp])
        json_code, json_out, json_err = _run([self.tmp, "--json", self.out])
        self.assertEqual(plain_out, json_out, "--json 改变了人读 stdout")
        self.assertEqual(plain_err, json_err, "--json 改变了 stderr")
        self.assertEqual(plain_code, json_code, "--json 改变了 exit code")

    def test_json_carries_the_facts_the_panel_needs(self):
        code, out, _ = _run([self.tmp, "--json", self.out])
        self.assertEqual(code, 0, out.decode("utf-8", "replace"))
        with io.open(self.out, encoding="utf-8") as f:
            data = json.loads(f.read())
        for key in ("tool", "verdict", "items"):
            self.assertIn(key, data, "既有字段不许丢")
        self.assertIn("facts", data, "面板要的事实层缺席")
        facts = data["facts"]
        self.assertEqual(facts.get("pipeline_status"), "gate3_approved")
        self.assertIsNotNone(facts.get("hours_since_last_turn"))
        self.assertEqual(facts.get("progress_first_line"), "对接窗口3 已过,正在写正本")
        stamps = facts.get("stamps") or {}
        self.assertEqual(set(stamps), {"task_plan.md", "dossier.md", "outline.md"})
        for name, stamp in stamps.items():
            self.assertEqual(stamp.get("status"), "approved", name)
            # hash 是面板要显示的东西之一,不能只在脚本里比完就扔
            self.assertRegex(str(stamp.get("current_hash")), r"^[0-9a-f]{12}$")
            self.assertEqual(stamp.get("approved_hash"), stamp.get("current_hash"))

    def test_facts_absent_for_tools_that_do_not_set_them(self):
        """Report 是共享类:没设事实的工具,JSON 必须与从前逐字节一致。"""
        from pipeline_lib import Report
        rep = Report("some_other_tool")
        rep.add("检查", "PASS", "细节")
        self.assertNotIn("facts", json.loads(rep.to_json()))




class StaleIsReportedWhenTheBodyMovedPastTheStamp(unittest.TestCase):
    """正本改过、印章没重盖的那一类项目:机器必须报 stale。

    ⚠️ 变异必须落在**参与 hash 的字节**上。规格里正文是 body.strip() 之后
    才进 hash,所以往末尾追加空白**不会**改变 hash —— 那是有意设计(避免尾部
    空白造成假 stale),不是缺陷。拿它当变异,得到的绿是假绿。
    所以下面先立守卫,再判产品。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-stale-")
        _project(self.tmp)
        self.out = os.path.join(self.tmp, "status.json")

    def _stamps(self, expect_code=0):
        """印章表。exit code 一并钉住:章一 stale,窗口就必须失败(1)。"""
        code, out, _ = _run([self.tmp, "--json", self.out])
        self.assertEqual(code, expect_code, out.decode("utf-8", "replace"))
        with io.open(self.out, encoding="utf-8") as f:
            return json.loads(f.read())["facts"]["stamps"]

    def _rewrite_dossier(self, tail):
        path = os.path.join(self.tmp, "dossier.md")
        meta, body = read_md(path)
        write_md(path, meta, body + tail)

    def test_trailing_whitespace_does_not_move_the_hash(self):
        """守卫本身的守卫:证明"尾部空白不算改动"确实是现行规格。

        这条一旦变红,说明规格改了,上面那段告诫和下面那条的选材都要重写。
        """
        before = self._stamps()["dossier.md"]["current_hash"]
        self._rewrite_dossier("   \n")
        self.assertEqual(self._stamps()["dossier.md"]["current_hash"], before)

    def test_an_edited_body_turns_that_one_stamp_stale(self):
        before = self._stamps()["dossier.md"]["current_hash"]
        self._rewrite_dossier("\n补一句可见的正文。\n")
        after = self._stamps(expect_code=1)
        # 守卫:先证明手真的碰到了被测量的那个量,再看它响不响。
        self.assertNotEqual(after["dossier.md"]["current_hash"], before,
                            "变异没落在 hash 上,下面的判定不成立")
        self.assertEqual(after["dossier.md"]["status"], "stale")
        # 反向:另两章不许被误伤,否则"变红"可能只是整体崩了。
        self.assertEqual(after["task_plan.md"]["status"], "approved")
        self.assertEqual(after["outline.md"]["status"], "approved")




class CountsComeFromOneImplementation(unittest.TestCase):
    """计数区的数:全部由真源算出,且口径只有一份。

    ⚠️ 这一族数踩过一次坑:cite_check 排 deprecated、render_html 不排,当前
    恰好相等只因为这批卡没有一张带该键。所以这里同时钉两件事 —— 数对不对,
    以及**排 deprecated 这个口径确实在生效**。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-counts-")
        _project(self.tmp)
        os.makedirs(os.path.join(self.tmp, "cards"))
        for i in range(3):
            card = os.path.join(self.tmp, "cards", "M%02d.md" % i)
            with io.open(card, "w", encoding="utf-8", newline="") as f:
                f.write("---" + NL + "id: M%02d" % i + NL + "title: 卡%d" % i + NL + "---" + NL + NL + "正文" + NL)
        self.out = os.path.join(self.tmp, "status.json")

    def _counts(self):
        _run([self.tmp, "--json", self.out])
        with io.open(self.out, encoding="utf-8") as f:
            return json.loads(f.read())["facts"]["counts"]

    def test_cards_are_counted_off_disk(self):
        self.assertEqual(self._counts()["cards_actual"], 3)

    def test_a_deprecated_card_leaves_the_count(self):
        """⭐ 口径判据:弃用一张,实测值必须减一。

        这条不红就说明「排 deprecated」没生效 —— 而那正是与 render_html
        分叉的那个键。
        """
        before = self._counts()["cards_actual"]
        path = os.path.join(self.tmp, "cards", "M01.md")
        text = io.open(path, encoding="utf-8").read()
        text = text.replace("title: 卡1", "title: 卡1" + NL + "deprecated: true")
        io.open(path, "w", encoding="utf-8", newline="").write(text)
        self.assertEqual(self._counts()["cards_actual"], before - 1)

    def test_declarations_come_from_the_products_too(self):
        """⭐ 自述不止 dossier 一路。

        ⚠️ 这条测的不是「有没有两份实现」,而是「同一份实现有没有被喂全」。
        `declared_card_counts` 一直是共用的,但 pipeline_status 曾经给它传空
        列表,于是遍历成品那一路一次都不跑 —— 面板只看见 dossier 的自述,
        交付对接窗口看见全部,同一屏两格给出相反读数。
        """
        out = os.path.join(self.tmp, "out")
        os.makedirs(out)
        with io.open(os.path.join(out, "成品.html"), "w", encoding="utf-8", newline="") as f:
            f.write("<html><body><p>本报告含 7 张证据卡。</p></body></html>")
        sources = [d["source"] for d in self._counts()["cards_declared"]]
        self.assertIn("成品.html", sources)

    def test_absent_sources_are_none_not_zero(self):
        """没有 references.yaml 时是 None,不是 0。

        0 是「数出来是零」,None 是「没能数」。面板对这两件事的说法不一样。
        """
        counts = self._counts()
        self.assertIsNone(counts["references_registered"])
        self.assertIsNone(counts["inline_citations"])




class CommitmentsAreParsedNotGuessed(unittest.TestCase):
    """承诺四组:两种书写形式、四个来源、三个状态。

    ⚠️ 盘上两种形式都在用(`1. **标签**:` 与 `- ① 文字`),
    第四组还独占一个没有锚点的文件。任何一条漏掉,面板会安静地少显示一组。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-commit-")
        _project(self.tmp)
        self.out = os.path.join(self.tmp, "status.json")

    def _groups(self):
        _run([self.tmp, "--json", self.out])
        with io.open(self.out, encoding="utf-8") as f:
            return {g["record"]: g for g in json.loads(f.read())["facts"]["commitments"]}

    def _write_section(self, record, lines):
        path = os.path.join(self.tmp, record)
        meta, body = read_md(path)
        write_md(path, meta, body + NL + "## 本轮承诺 {#commitments}" + NL + NL + NL.join(lines) + NL)

    def test_numbered_form_is_read(self):
        self._write_section("dossier.md", [
            "1. **口径承诺**:没有变更。",
            "2. **最脆证据**:M06 基准。",
            "3. **最可能错在哪**:低估同址竞争。",
        ])
        group = self._groups()["dossier.md"]
        self.assertEqual(group["state"], "ok")
        self.assertEqual(len(group["items"]), 3)
        self.assertEqual(group["items"][0]["label"], "口径承诺")

    def test_circled_form_is_read(self):
        self._write_section("outline.md", [
            "- ① 本轮我改了什么:章节顺序。",
            "- ② 本结论最脆的一处证据:折扣量级。",
            "- ③ 如果我错了,最可能错在哪:优先级。",
        ])
        group = self._groups()["outline.md"]
        self.assertEqual(group["state"], "ok")
        self.assertEqual(len(group["items"]), 3)

    def test_only_the_commitments_section_is_read(self):
        """⭐ 正文别处的有序列表不算数。

        不加这个边界,实测某份 task_plan 会数出 12 条而实际是 3 条。
        """
        path = os.path.join(self.tmp, "dossier.md")
        meta, body = read_md(path)
        write_md(path, meta, body + NL + "## 别的一节" + NL + NL
                 + "1. **看起来像**:但不在承诺段里。" + NL + NL
                 + "## 本轮承诺 {#commitments}" + NL + NL
                 + "1. **口径承诺**:唯一真条目。" + NL + NL
                 + "## 之后的一节" + NL + NL
                 + "2. **也不算**:在承诺段之后。" + NL)
        items = self._groups()["dossier.md"]["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["label"], "口径承诺")

    def test_the_delivery_group_has_no_anchor_and_is_still_found(self):
        """第四组独占一个文件,整篇就是承诺块,没有 {#commitments}。"""
        lib = os.path.join(self.tmp, "library")
        os.makedirs(lib)
        with io.open(os.path.join(lib, "delivery_commitments.md"), "w",
                     encoding="utf-8", newline="") as f:
            f.write("# 交付判断三行" + NL + NL
                    + "- ① 改了什么:无。" + NL
                    + "- ② 最脆的一处:基准。" + NL
                    + "- ③ 最可能错在哪:上限。" + NL)
        group = self._groups()["delivery"]
        self.assertEqual(group["state"], "ok")
        self.assertEqual(len(group["items"]), 3)

    def test_a_project_after_the_convention_without_commitments_is_a_defect(self):
        """⭐ 这一格存在的全部理由。

        判别式是交付日期,不是段落缺失 —— 用后者的话,将来真漏写承诺的新项目
        会被伪装成历史遗留,而那正是要抓的东西。盘上还没有这个例子,所以它
        只能在这里被造出来。

        ⚠️ 日期必须显式设:公用 fixture 的批准日是 2020 年,落在约定生效之前,
        不改的话这条测试会以「早于约定」通过 —— 一条永远不会红的哨兵。
        """
        path = os.path.join(self.tmp, "task_plan.md")
        meta, body = read_md(path)
        meta["approval"]["approved_at"] = "2026-08-29T16:59:45-04:00"
        write_md(path, meta, body)
        for group in self._groups().values():
            self.assertEqual(group["state"], "missing", group["record"])

    def test_a_project_delivered_before_the_convention_is_not_a_defect(self):
        path = os.path.join(self.tmp, "task_plan.md")
        meta, body = read_md(path)
        meta["approval"]["approved_at"] = "2026-08-24T17:01:19-04:00"
        write_md(path, meta, body)
        for group in self._groups().values():
            self.assertEqual(group["state"], "predates", group["record"])




class LinksAlignAcrossFormats(unittest.TestCase):
    """一条带 & 的 URL,在 html / docx / 登记表三处必须是同一个字符串。

    ⚠️ 两个格式各自转义 &:HTML 写 &amp;,docx 的 rels 也写 &amp;,而
    references.yaml 存的是真实 URL。检查器 ③「文内⊆参考」拿转义后的形态去
    比对原形,11 条实际登记过的链接被判成没登记 —— 实测。
    ⭐ 这条测试的意义不在「反转义能跑」,在「三处对齐」:任何一处漏掉,
    集合比较就会静默地把登记过的链接算成幽灵。
    """

    URL = "https://example.com/corp/go.php?vt=4&stockid=100001"
    ESCAPED = "https://example.com/corp/go.php?vt=4&amp;stockid=100001"

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-amp-")

    def _html(self):
        p = os.path.join(self.tmp, "成品.html")
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<html><body><p>见 <a href=" + chr(34) + self.ESCAPED + chr(34)
                    + ">来源</a></p></body></html>")
        return p

    def _docx(self):
        import zipfile
        p = os.path.join(self.tmp, "成品.docx")
        rels = ('<?xml version="1.0"?><Relationships>'
                '<Relationship Id="rId1" Type="hyperlink" Target="' + self.ESCAPED
                + '" TargetMode="External"/></Relationships>')
        doc = ('<?xml version="1.0"?><w:document xmlns:w="w" xmlns:r="r">'
               '<w:hyperlink r:id="rId1"><w:t>来源</w:t></w:hyperlink></w:document>')
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/_rels/document.xml.rels", rels)
            z.writestr("word/document.xml", doc)
        return p

    def test_all_three_forms_agree(self):
        from pipeline_lib import docx_links, html_links
        self.assertEqual(html_links(self._html()), [self.URL])
        self.assertEqual(docx_links(self._docx()), [self.URL])

    def test_the_numeric_entity_form_also_aligns(self):
        """`&#38;` 与 `&amp;` 是两条路径。命名实体过了不等于数字实体也过。"""
        from pipeline_lib import html_links
        p = os.path.join(self.tmp, "num.html")
        url = self.URL.replace("&", "&#38;")
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<a href=" + chr(34) + url + chr(34) + ">来源</a>")
        self.assertEqual(html_links(p), [self.URL])

    def test_percent_encoding_is_left_alone(self):
        """⛔ 反过来的一半:百分号编码**不是**实体,不该被反转义。

        `%26` 在 URL 里是合法字符序列,把它变成 `&` 会造出一个不同的 URL ——
        那样登记表反而对不上。这条防的是「归一化做得太多」。
        """
        from pipeline_lib import html_links
        p = os.path.join(self.tmp, "pct.html")
        url = "https://example.com/a?x=1%26y=2"
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<a href=" + chr(34) + url + chr(34) + ">来源</a>")
        self.assertEqual(html_links(p), [url])

    def test_the_escaped_form_never_reaches_a_caller(self):
        """⛔ 反向:转义形态一个字符都不许流出去。

        只断言「等于真实 URL」还不够 —— 如果将来有人改成半反转义(比如只处理
        &amp; 不处理 &#38;),前一条仍可能通过而这条会红。
        """
        from pipeline_lib import docx_links, html_links
        for got in (html_links(self._html()), docx_links(self._docx())):
            self.assertNotIn("&amp;", got[0])
            self.assertIn("&stockid=", got[0])



class TablesAreOneCaliber(unittest.TestCase):
    """「有几张表」只有一份判定,两个渲染器共用。

    ⚠️ 这一族踩过的坑:docx 渲染器根本没有表格代码,markdown 表被当段落写进
    成品 —— 实测真表格 0、残留竖线 233、`**` 18,而 HTML 那边 6 张都是
    好的。交付对接窗口 ④ 只比链接集合,看不见这种退化。
    """

    def test_a_table_needs_a_rule_line_under_its_head(self):
        from pipeline_lib import count_tables, split_table
        lines = ["| 甲 | 乙 |", "|---|---|", "| 1 | 2 |", "", "普通段落"]
        head, rows, nxt = split_table(lines, 0)
        self.assertEqual(head, ["甲", "乙"])
        self.assertEqual(rows, [["1", "2"]])
        self.assertEqual(nxt, 3)
        self.assertEqual(count_tables(chr(10).join(lines)), 1)

    def test_a_pipe_line_without_a_rule_is_not_a_table(self):
        """⛔ 判据是「下一行是分隔行」,不是「本行以竖线开头」。

        正文里以竖线开头的普通行、以及表格续行,都不能被当成新表的开始 ——
        否则一张表会被数成好几张,而 ⑫ 正要靠这个数判双格式是否一致。
        """
        from pipeline_lib import count_tables, split_table
        self.assertIsNone(split_table(["| 这只是一行文字", "后面没有分隔行"], 0))
        self.assertEqual(count_tables("| 甲 | 乙 |" + chr(10) + "| 丙 | 丁 |"), 0)

    def test_breaking_the_rule_line_drops_the_count(self):
        """对照组:分隔行一坏,这张表就不算表 —— 判据能分辨真假。"""
        from pipeline_lib import count_tables
        good = "| 甲 | 乙 |" + chr(10) + "|---|---|" + chr(10) + "| 1 | 2 |"
        self.assertEqual(count_tables(good), 1)
        self.assertEqual(count_tables(good.replace("---", "===")), 0)


class DocxCarriesRealTables(unittest.TestCase):
    """docx 成品里的表必须是 w:tbl,不是一段带竖线的文字。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-docx-")

    def _render(self, body):
        draft = os.path.join(self.tmp, "稿.md")
        refs = os.path.join(self.tmp, "references.yaml")
        out = os.path.join(self.tmp, "成品.docx")
        with io.open(draft, "w", encoding="utf-8", newline="") as f:
            f.write("---" + NL + "title: 测试稿" + NL + "---" + NL + NL + body + NL)
        with io.open(refs, "w", encoding="utf-8", newline="") as f:
            f.write("[]" + NL)
        import subprocess
        r = subprocess.run([sys.executable, "-X", "utf8",
                            os.path.join(HERE, "render_docx.py"), draft, refs, out],
                           capture_output=True)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        import docx
        d = docx.Document(out)
        text = NL.join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    text += NL + c.text
        return d, text

    def test_a_markdown_table_becomes_a_real_table(self):
        d, text = self._render("| 甲 | 乙 |" + NL + "|---|---|" + NL + "| 1 | 2 |")
        self.assertEqual(len(d.tables), 1)
        # ⛔ 残留竖线为 0:成品里的竖线就是「这张表没被渲染」的可见证据。
        self.assertEqual(text.count("|"), 0)
        self.assertIn("甲", d.tables[0].rows[0].cells[0].text)

    def test_bold_markers_do_not_survive_into_the_product(self):
        """`**` 是标记不是内容。它出现在成品里就是渲染没做完。"""
        d, text = self._render("这里有**重点**。" + NL + NL
                               + "| **表头** | 乙 |" + NL + "|---|---|" + NL + "| **粗** | 2 |")
        self.assertEqual(text.count("**"), 0)
        self.assertIn("重点", text)
        self.assertIn("粗", text)

    def test_a_table_glued_to_a_paragraph_is_still_found(self):
        """表可以紧贴在正文行之后 —— 块之间才有空行,不能假设一块只有一种东西。"""
        d, _ = self._render("先说一句。" + NL + "| 甲 | 乙 |" + NL + "|---|---|" + NL + "| 1 | 2 |")
        self.assertEqual(len(d.tables), 1)



class ResidueIsFoundWhereverItSits(unittest.TestCase):
    """markdown 残留判据:位置无关、实体解过、不误伤正常文字。

    ⚠️ 第一版用 `^...$` 锚定行首行尾,在一份明显坏掉的 docx 上报了
    rule=0 —— 因为 docx 的取文把整张表压成一行,分隔行夹在文字中间。
    ⭐ 漏报比漏一项更糟:窗口会说「干净」。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-res-")

    def test_a_rule_row_is_found_mid_line(self):
        from pipeline_lib import markdown_residue
        collapsed = "| 甲 | 乙 || --- | --- || 1 | 2 |"
        self.assertEqual(markdown_residue(collapsed)["rule"], 1)

    def test_entity_encoded_pipes_are_found_after_unescaping(self):
        """HTML 里一根竖线可以写成 `&#124;`。不解实体就数不到 —— 而数不到
        在这里等于「报告说干净」。`html_text` 负责解,判据负责数。
        """
        from pipeline_lib import html_text, markdown_residue
        p = os.path.join(self.tmp, "x.html")
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<html><body><p>&#124; 甲 &#124; 乙 &#124;&#124; --- &#124; --- &#124;</p></body></html>")
        self.assertEqual(markdown_residue(html_text(p))["rule"], 1)

    def test_named_entities_are_unescaped_too(self):
        """数字实体与命名实体是两条路径。`&#124;` 过了不等于 `&verbar;` 也过。"""
        from pipeline_lib import html_text, markdown_residue
        p = os.path.join(self.tmp, "named.html")
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<html><body><p>&verbar; 甲 &verbar; 乙 &verbar;&verbar; --- &verbar; --- &verbar;</p>"
                    "<p>&ast;&ast;粗&ast;&ast;</p></body></html>")
        res = markdown_residue(html_text(p))
        self.assertEqual(res["rule"], 1)
        self.assertEqual(res["bold"], 2)

    def test_a_sharp_after_a_letter_is_not_a_heading(self):
        """⛔ 不许误伤:`C# 语言` 不是标题标记。误报会让干净成品被判 FAIL。"""
        from pipeline_lib import markdown_residue
        self.assertEqual(markdown_residue("用 C# 语言写的")["heading"], 0)
        self.assertEqual(markdown_residue("正文。## 二、下一节")["heading"], 1)

    def test_clean_text_is_clean(self):
        from pipeline_lib import markdown_residue
        self.assertEqual(markdown_residue("一段普通正文,没有任何标记。"),
                         {"rule": 0, "bold": 0, "heading": 0})


class EachFormatIsCountedOnItsOwn(unittest.TestCase):
    """真表格按格式各自数,口径与正本源同一份。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-fmt-")

    def test_html_tables_are_counted(self):
        from pipeline_lib import count_tables_html
        p = os.path.join(self.tmp, "x.html")
        with io.open(p, "w", encoding="utf-8", newline="") as f:
            f.write("<div><table><tr><td>1</td></tr></table><TABLE></TABLE></div>")
        self.assertEqual(count_tables_html(p), 2)

    def test_docx_tables_are_counted_without_python_docx(self):
        """⭐ 用 zipfile + ElementTree 数,不引入 python-docx:交付对接窗口跑不起来
        比交付对接窗口少一项更糟,不该为了数表格给它加运行时依赖。
        """
        import zipfile
        from pipeline_lib import count_tables_docx
        p = os.path.join(self.tmp, "x.docx")
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        doc = ('<?xml version="1.0"?><w:document xmlns:w="' + ns + '">'
               '<w:tbl><w:tr><w:tc><w:p><w:t>1</w:t></w:p></w:tc></w:tr></w:tbl>'
               '<w:p><w:t>中间段落</w:t></w:p>'
               '<w:tbl><w:tr><w:tc><w:p><w:t>2</w:t></w:p></w:tc></w:tr></w:tbl>'
               '</w:document>')
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", doc)
        self.assertEqual(count_tables_docx(p), 2)



class BomIsNotContent(unittest.TestCase):
    r"""v1.0.4:BOM 不算内容 —— 同一份内容加不加它,印章值必须一样。

    ⚠️ 它同时坏三件事:frontmatter 的 `\A---` 匹配不上、标题的 `#` 匹配不上、
    而 `str.strip()` 不去 U+FEFF 所以它进了 content_hash。实测:整稿
    首行是 BOM + `# 标题`,交付 HTML 里留着一个字面的 `#`。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-bom-")
        self.body = ("---" + NL + "title: 甲" + NL + "version: 1" + NL + "---"
                     + NL + NL + "正文一句。" + NL)

    def _write(self, name, text):
        p = os.path.join(self.tmp, name)
        io.open(p, "w", encoding="utf-8", newline="").write(text)
        return p

    def test_the_hash_is_the_same_with_or_without_a_bom(self):
        from pipeline_lib import content_hash, read_md
        plain = read_md(self._write("a.md", self.body))
        withbom = read_md(self._write("b.md", chr(65279) + self.body))
        self.assertEqual(content_hash(*plain), content_hash(*withbom))

    def test_one_more_real_character_does_change_the_hash(self):
        """反向:判据得能分辨真假 —— 只证明「BOM 不改 hash」不够,
        一个把什么都当无关的实现也能通过那条。
        """
        from pipeline_lib import content_hash, read_md
        plain = read_md(self._write("a.md", self.body))
        longer = read_md(self._write("c.md", self.body.replace("正文一句。", "正文一句。x")))
        self.assertNotEqual(content_hash(*plain), content_hash(*longer))

    def test_a_bom_no_longer_hides_the_frontmatter(self):
        """BOM 挡在 `---` 前面时,frontmatter 会被整块当成正文。"""
        from pipeline_lib import read_md
        meta, body = read_md(self._write("d.md", chr(65279) + self.body))
        self.assertEqual(meta.get("title"), "甲")
        self.assertNotIn("---", body)



class TheDocumentTitleIsNotAChapter(unittest.TestCase):
    """稿件首行的一级标题 = 文档名,不是第 1 章。

    ⚠️ 实测:BOM 挡住标题判定,那行被渲成正文段落 —— 交付 HTML 里留着
    一个字面的 `#`。剥掉 BOM 之后它成了正式章节,章节 20→21、目录多一条、
    锚点全位移。两个都不对:它和 --title 是同一件事的两种写法。
    """

    def test_a_leading_h1_is_taken_as_the_title(self):
        from pipeline_lib import take_doc_title
        title, rest = take_doc_title("# 文档名" + NL + NL + "## 一、章" + NL + "正文")
        self.assertEqual(title, "文档名")
        self.assertTrue(rest.startswith("## 一、章"))

    def test_a_leading_h2_is_left_alone(self):
        """⛔ 只认一级:以 `##` 开头的稿件一个字都不许动。"""
        from pipeline_lib import take_doc_title
        body = "## 一、章" + NL + "正文"
        self.assertEqual(take_doc_title(body), (None, body))

    def test_an_h1_further_down_is_not_the_title(self):
        """⛔ 只认最前面那个:正文中间的一级标题是章节,摘走它会丢内容。"""
        from pipeline_lib import take_doc_title
        body = "## 一、章" + NL + NL + "# 不该被摘"
        self.assertEqual(take_doc_title(body), (None, body))

    def test_the_html_renderer_actually_drops_it(self):
        """⭐ 接线也要测,不只测函数。

        ⚠️ 我把 `take_doc_title` 的导入写进了一个跨行 import 的括号中间,
        整个渲染器语法错 —— 单元测试全绿,因为它们不 import 渲染器。
        「函数对」和「函数被正确接上」是两件事。
        """
        import glob
        import subprocess
        tmp = tempfile.mkdtemp(prefix="pls-wire-")
        draft = os.path.join(tmp, "稿.md")
        refs = os.path.join(tmp, "references.yaml")
        out = os.path.join(tmp, "成品.html")
        io.open(draft, "w", encoding="utf-8", newline="").write(
            chr(65279) + "# 文档名" + NL + NL + "## 一、第一章" + NL + NL + "正文。" + NL)
        io.open(refs, "w", encoding="utf-8", newline="").write("[]" + NL)
        r = subprocess.run([sys.executable, "-X", "utf8",
                            os.path.join(HERE, "render_html.py"), draft, refs, out],
                           capture_output=True)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        html = io.open(out, encoding="utf-8").read()
        # 文档名不出现在目录里,也不作为杂段落留在正文里
        self.assertNotIn("<p>文档名</p>", html)
        self.assertIn("第一章", html)
        # 章节数按真章节算,不含文档名
        self.assertIn(">1</span>第一章", html.replace("一、", ""))
        # ⛔ BOM 一个都不许流进成品
        self.assertEqual(html.count(chr(65279)), 0)

    def test_both_renderers_agree(self):
        """⭐ 判定共用一份:各写一份就会出现「HTML 少一行、DOCX 没少」这类
        只在某个格式里可见的差异,而 ④ 只比链接集合看不见它。
        """
        import render_docx
        import render_html
        from pipeline_lib import take_doc_title
        for mod in (render_html, render_docx):
            self.assertIs(mod.take_doc_title, take_doc_title)



class ReGatingMustBumpTheVersion(unittest.TestCase):
    """重过窗口要先升版再落章,否则窗口报告与正本对不上。

    ⚠️ 实测:dossier 的 revision_log 三条都写 v=2 而 frontmatter 仍
    version=1,对接窗口2 报告停在 v1/hash 0032c6/157 来源,而 dossier 已到
    0c55fd/164 条 —— 报告里的数字是旧版的,没人会发现。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-ver-")

    def _record(self, version, log_versions):
        p = os.path.join(self.tmp, "dossier.md")
        meta = {"title": "证据汇编", "version": version,
                "approval": {"status": "awaiting"},
                "revision_log": [{"v": v, "why": "改了点东西"} for v in log_versions]}
        write_md(p, meta, NL + "正文若干。" + NL)
        return p

    def _approve(self, path):
        return _run_stamp([path, "--approve", "--by", "reviewer", "--quote", "确认通过"])

    def test_a_log_ahead_of_the_version_is_refused(self):
        code, out, _ = self._approve(self._record(1, [2, 2, 2]))
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))
        self.assertIn("改了正本要先升版", out.decode("utf-8", "replace"))

    def test_one_entry_ahead_is_refused_too(self):
        """⛔ 判据不是「条数 > version」:另一份 dossier 只有一条 v=2 却停在
        version=1,条数比会放它过去。真正的问题是**日志声称的版本**超过了
        frontmatter 到达的版本,与条数无关。
        """
        code, out, _ = self._approve(self._record(1, [2]))
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))

    def test_many_entries_within_one_version_are_fine(self):
        """⛔ 反向:同一版内改三次是正常的,会留三条 v 相同的记录。
        用「条数 > version」比会误伤这种 —— 而误伤比漏报更快让人把守卫关掉。
        """
        code, out, _ = self._approve(self._record(2, [2, 2, 2]))
        self.assertEqual(code, 0, out.decode("utf-8", "replace"))

    def test_a_record_without_a_log_is_fine(self):
        code, out, _ = self._approve(self._record(1, []))
        self.assertEqual(code, 0, out.decode("utf-8", "replace"))

    def _bytes(self, path):
        return io.open(path, "rb").read()

    def test_refusing_does_not_touch_the_file(self):
        """⛔ 守卫确实在 write_md 之前 return,但没有一条测试钉住这一点 ——
        把写盘挪到守卫之前,上面四条照样全绿(实测验证过)。
        比的是字节而不是 md5:同样的判据,少一个 import,且更强。
        """
        p = self._record(1, [2])
        before = self._bytes(p)
        code, out, _ = self._approve(p)
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))
        self.assertEqual(self._bytes(p), before, "拒绝落章却动了盘上的字节")

    def test_a_string_version_is_refused_not_skipped(self):
        """`version: '1'` —— 原来 isinstance 过滤把守卫整条跳过,rc=0 落章。"""
        p = self._record("1", [2])
        before = self._bytes(p)
        code, out, _ = self._approve(p)
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))
        self.assertEqual(self._bytes(p), before)

    def test_a_missing_version_is_refused(self):
        p = os.path.join(self.tmp, "no-version.md")
        write_md(p, {"title": "证据汇编", "approval": {"status": "awaiting"},
                     "revision_log": [{"v": 2, "why": "改了点东西"}]},
                 NL + "正文若干。" + NL)
        before = self._bytes(p)
        code, out, _ = self._approve(p)
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))
        self.assertEqual(self._bytes(p), before)

    def test_a_string_v_in_the_log_is_refused(self):
        """`v: '2'` —— 原来被 isinstance 过滤掉,claimed 变空,守卫等于没有。"""
        p = self._record(1, ["2"])
        before = self._bytes(p)
        code, out, _ = self._approve(p)
        self.assertEqual(code, 2, out.decode("utf-8", "replace"))
        self.assertEqual(self._bytes(p), before)

    def _write_raw_log(self, version, log_yaml):
        """直接拼 frontmatter:write_md 只会写出合法的列表,而这里要的正是**不合法**的形状。"""
        p = os.path.join(self.tmp, "raw.md")
        text = ("---" + NL + "title: 汇编" + NL + "version: %d" % version + NL
                + "approval:" + NL + "  status: awaiting" + NL
                + log_yaml + NL + "---" + NL + NL + "正文若干。" + NL)
        with io.open(p, "w", encoding="utf-8", newline=NL) as f:
            f.write(text)
        return p

    def test_invalidate_refuses_a_non_integer_version_instead_of_crashing(self):
        """⛔ `--invalidate` 用的是裸 `int(meta.get("version") or 1)`,同一份正本
        在 `--approve` 上被规规矩矩拒批,在这条路上抛 ValueError 崩掉。判据长在
        一条路径里,另外两条各自用裸转换 —— 现在三条共用模块级 `_is_int`。
        (对抗核查 2026-09-04。)"""
        p = self._write_raw_log(1, "revision_log: []")
        text = io.open(p, encoding="utf-8").read().replace("version: 1", "version: 'v2'")
        with io.open(p, "w", encoding="utf-8", newline=NL) as f:
            f.write(text)
        before = io.open(p, "rb").read()
        code, out, err = _run_stamp([p, "--invalidate", "--why", "改口径"])
        blob = (out + err).decode("utf-8", "replace")
        self.assertEqual(code, 2, blob)
        self.assertNotIn("Traceback", blob)
        self.assertEqual(io.open(p, "rb").read(), before)

    def test_advance_without_a_state_is_a_message_not_an_indexerror(self):
        """⛔ `argv[argv.index("--advance") + 1]` 是裸下标:少写状态名就是
        IndexError + traceback,而**紧邻的**拼错状态名那一支处理得好好的。"""
        p = self._write_raw_log(1, "revision_log: []")
        code, out, err = _run_stamp([p, "--advance"])
        blob = (out + err).decode("utf-8", "replace")
        self.assertEqual(code, 2, blob)
        self.assertNotIn("IndexError", blob)

    def test_a_revision_log_that_is_not_a_list_of_dicts_is_refused(self):
        """⛔ 少一个 `-` 就落章。

        原来这里是个**过滤器**:`[e for e in log if isinstance(e, dict)]`。把整段
        写成 YAML 映射(少个 `-`)、写成散文、写成嵌套列表,过滤后 entries 全空,
        下面两道检查都空转,rc=0 落章 —— 而同一条 v2 声明带上 `-` 是拒批的。
        实测过:带 `-` rc=2 不写盘,少个 `-` rc=0 写盘且 `--check` 说「印章有效」。
        (对抗核查 2026-09-04。)
        """
        shapes = {
            "映射(少一个 -)": "revision_log:" + NL + "  v: 2" + NL + "  what: 改了",
            "散文条目": "revision_log:" + NL + '- "v2 2026-09-04 补了来源"',
            "标量字符串": 'revision_log: "v9 大改了"',
            "嵌套列表": "revision_log:" + NL + "- - v: 9",
            "整数(不可迭代)": "revision_log: 5",
        }
        for name, yaml in shapes.items():
            p = self._write_raw_log(1, yaml)
            before = io.open(p, "rb").read()
            code, out, _ = self._approve(p)
            self.assertEqual(code, 2, "%s 竟然落章了:%s" % (name, out.decode("utf-8", "replace")))
            self.assertEqual(io.open(p, "rb").read(), before, "%s:拒批却动了盘上的字节" % name)

    def test_a_boolean_version_is_refused(self):
        """isinstance(True, int) 为真 —— 不排除 bool 的话 True 会被当版本号用。

        ⛔ 样本必须让「日志超前」那条分支**够不着**,否则这条测试证明不了
        bool 排除:最初写的是 (True, [2]),而 True == 1、max 2 > 1,于是把
        `_is_int` 改成接受 bool 之后 9 条照绿 —— 它一直在测另一条分支。
        (对抗式独立复核抓到;在副本上复现了变异体存活。)
        这里两个样本都让 max(v) <= version 成立,拒批只能来自 bool 排除。
        """
        for log in ([1], []):
            p = self._record(True, log)
            before = self._bytes(p)
            code, out, _ = self._approve(p)
            self.assertEqual(code, 2, "version=True log=%s 竟然落章了:%s"
                             % (log, out.decode("utf-8", "replace")))
            self.assertEqual(self._bytes(p), before,
                             "version=True log=%s:拒批却动了盘上的字节" % (log,))



class SessionFactsReadsTheLogNotTheSelfReport(unittest.TestCase):
    """步数与项目外读取:从会话日志取,口径写死。

    ⚠️ 规格 v1.0.3 曾记「管线读不到 DSH 步数事件」——不成立。真实情况是
    没人去读:实测日志里 255 个 step/start,而 dossier 写着 null、正文还
    写了「DSH 环境无步数事件」。**把"我没去取"写成了"它不存在"。**
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-sess-")
        self.project = os.path.join(self.tmp, "proj")
        os.makedirs(self.project)
        self.home = os.path.join(self.tmp, "home")
        os.makedirs(os.path.join(self.home, ".dsh", "sessions", "tag", "session-x"))

    def _call(self, name, args):
        return {"type": "tool/call", "data": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}

    def _log(self, events):
        """写一份 multi-frame zstd 会话日志(用 node 压,跟读它的那一侧同源)。"""
        raw = os.path.join(self.tmp, "raw.jsonl")
        with io.open(raw, "w", encoding="utf-8", newline="") as f:
            for e in events:
                f.write(json.dumps(e, ensure_ascii=False) + NL)
        out = os.path.join(self.home, ".dsh", "sessions", "tag", "session-x", "session.jsonl.zstd")
        # ⚠️ node -e 的 argv:[execPath, ...`--` 之后的参数],没有脚本名那一位。
        script = ("const fs=require('fs'),z=require('zlib');"
                  "fs.writeFileSync(process.argv[2], z.zstdCompressSync(fs.readFileSync(process.argv[1])))")
        r = subprocess.run(["node", "-e", script, "--", raw, out], capture_output=True)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        return out

    def _facts(self, events):
        self._log(events)
        out = os.path.join(self.tmp, "facts.json")
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home)
        r = subprocess.run(["node", os.path.join(HERE, "session_facts.mjs"),
                            "--project", self.project, "--json", out],
                           capture_output=True, env=env)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        with io.open(out, encoding="utf-8") as f:
            return json.loads(f.read())

    def test_steps_come_from_the_log(self):
        inside = os.path.join(self.project, "task_plan.md")
        events = [{"type": "step/start", "data": {}} for _ in range(7)]
        events.append(self._call("read", {"file_path": inside}))
        self.assertEqual(self._facts(events)["steps_used"], 7)

    def test_a_read_outside_the_project_is_listed(self):
        inside = os.path.join(self.project, "task_plan.md")
        outside = os.path.join(self.tmp, "别处", "聊天记录.md")
        facts = self._facts([self._call("read", {"file_path": inside}),
                             self._call("read", {"file_path": outside})])
        paths = [r["path"] for r in facts["outside_reads"]]
        self.assertEqual(len(paths), 1)
        self.assertIn("聊天记录.md", paths[0])

    def test_a_grep_pattern_is_not_a_path(self):
        """⛔ grep 的 `pattern` 是正则。收它会把「单品|产能竞争」报成越界读取,
        而一段满是噪声的清单等于没有这一段 —— 第一版就是这么错的,42 条里 18 条是噪声。
        """
        inside = os.path.join(self.project, "task_plan.md")
        # ⛔ 样本必须是一条**绝对路径样子**的正则:原来用「单品|产能竞争|转向」,
        # 它本来就不是绝对路径,所以即使把 pattern 重新当路径收进来,清单也还是
        # 空的 —— 断言恒真,守不住那条回归(对抗核查 2026-09-04 指出)。
        bait = os.path.join(self.tmp, "桌面", "不该出现.md")
        facts = self._facts([self._call("read", {"file_path": inside}),
                             self._call("grep", {"pattern": bait})])
        self.assertEqual(facts["outside_reads"], [],
                         "grep 的 pattern 被当成路径收了进来")

    def test_a_relative_path_is_not_guessed_at(self):
        """⛔ 相对于哪个 cwd 我们不知道。拿进程 cwd 去 resolve 会造出根本
        没被读过的路径 —— 宁可漏,不可编。
        """
        inside = os.path.join(self.project, "task_plan.md")
        facts = self._facts([self._call("read", {"file_path": inside}),
                             self._call("glob", {"path": "projects/*"})])
        self.assertEqual(facts["outside_reads"], [])

    def test_no_matching_log_says_so_instead_of_zero(self):
        """⛔ 取不到时 steps_used 必须是 null 不是 0,而且**退出码非 0**。
        0 是「数出来是零步」,null 是「没能数」。
        ⚠️ 2026-09-03 契约收紧:原来这条路 rc=0,于是调用方一律当成功,空清单
        被读成「没越界」—— 正是四份 skill 明令禁止的那种沉默。
        """
        self._log([{"type": "step/start", "data": {}}])
        out = os.path.join(self.tmp, "facts.json")
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home)
        env.pop("DSH_SESSION_ID", None)
        env.pop("DSH_SESSION_JSONL", None)
        r = subprocess.run(["node", os.path.join(HERE, "session_facts.mjs"),
                            "--project", self.project, "--json", out],
                           capture_output=True, env=env)
        self.assertNotEqual(r.returncode, 0, "取不到会话却报了成功")
        with io.open(out, encoding="utf-8") as f:
            facts = json.loads(f.read())
        self.assertIsNone(facts["steps_used"])
        self.assertIn("不要写 0", facts["note"])


def _cite_project(tmp, steps_used, facts=None, dashboard_extra=None):
    """够 cite_check 跑到 ⑧ 的最小项目。"""
    dashboard = {"cards_count": 0, "steps_used": steps_used}
    facts_blob = None
    if facts is not None:
        if isinstance(facts, dict):
            facts_blob = dict(facts)
        else:
            facts_blob = {
                "steps_used": facts,
                "session": "session-test",
                "session_picked_by": "DSH_SESSION_ID(harness 注入)",
                "captured_at": "2026-09-04T12:00:00.000Z",
            }
            dashboard.update({k: facts_blob[k]
                              for k in ("session", "session_picked_by", "captured_at")})
    dashboard.update(dashboard_extra or {})
    write_md(os.path.join(tmp, "task_plan.md"),
             {"title": "计划", "version": 1, "budget": {"steps_max": 500}},
             NL + "正文。" + NL)
    write_md(os.path.join(tmp, "dossier.md"),
             {"title": "汇编", "version": 1,
              "dashboard": dashboard},
             NL + "正文。" + NL)
    with io.open(os.path.join(tmp, "references.yaml"), "w", encoding="utf-8", newline=NL) as f:
        f.write("refs: []" + NL)
    for d in ("drafts", "out", "cards"):
        os.makedirs(os.path.join(tmp, d))
    if facts_blob is not None:
        with io.open(os.path.join(tmp, "out", "session_facts.json"), "w",
                     encoding="utf-8", newline=NL) as f:
            f.write(json.dumps(facts_blob))
    return tmp


def _run_cite(project):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    r = subprocess.run([sys.executable, "-X", "utf8", os.path.join(HERE, "cite_check.py"),
                        "--project", project, "--skip-stamps"], capture_output=True, env=env)
    return (r.stdout + r.stderr).decode("utf-8", "replace")


def _steps_line(text):
    """⛔ 按**项名**取,不按「步数」两个字取:预算仪表那一行的修复提示里也有
    「步数跑 session_facts.mjs 取」,而它排在前面 —— 太泛的子串会匹配到
    另一项的提示语,断言就落在了别的东西上(对抗核查 2026-09-04 点出的形态)。
    """
    for line in text.split(NL):
        if "预算仪表·步数:" in line:
            return line.strip()
    return "(输出里没有预算仪表·步数这一项)"


class StepsUsedNeedsAMachineSource(unittest.TestCase):
    """⑧ 的 steps_used 分支 —— 技能里那条 ⛔「不许自报」的机器兜底。

    ⛔ 原来 `dash.get("steps_used", 0)` 把 null 当 0:超限判断永不触发,而输出
    还写着「预算内」。⛔ 而 v1.0.3 的注释「步数只存在于 DSH session 事件里,
    脚本读不到」是**错的** —— 实测日志里有 255 个 `step/start`。那句假话让
    这项检查十天里一直在替自报数背书。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cc-steps-")

    def test_null_is_not_taken_not_zero(self):
        _cite_project(self.tmp, None)
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("未取到", line)

    def test_a_number_with_no_machine_source_is_named_self_reported(self):
        _cite_project(self.tmp, 30)
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("自报", line)

    def test_disagreeing_with_session_facts_is_a_fail(self):
        _cite_project(self.tmp, 30, facts=255)
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("FAIL", line)
        self.assertIn("255", line)

    def test_a_missing_key_is_also_not_zero(self):
        """⛔ 上一条用的是 `steps_used: null`,而 `dict.get(k, 0)` 的默认值只在
        **键缺失**时生效 —— 所以那条测试根本盖不住它要防的那个默认值。这条补上
        键整个不在的形状。(对抗核查 2026-09-04 指出。)"""
        write_md(os.path.join(self.tmp, "task_plan.md"),
                 {"title": "计划", "version": 1, "budget": {"steps_max": 500}},
                 NL + "正文。" + NL)
        write_md(os.path.join(self.tmp, "dossier.md"),
                 {"title": "汇编", "version": 1, "dashboard": {"cards_count": 0}},
                 NL + "正文。" + NL)
        with io.open(os.path.join(self.tmp, "references.yaml"), "w",
                     encoding="utf-8", newline=NL) as f:
            f.write("refs: []" + NL)
        for d in ("drafts", "out", "cards"):
            os.makedirs(os.path.join(self.tmp, d))
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("未取到", line)

    def test_a_json_that_is_there_but_holds_null_is_not_reported_as_absent(self):
        """⛔ 「文件不在」会让人去跑脚本,而真正的毛病可能是脚本跑了却没取到。
        三种原因(不在 / 读不出 / 里面是 null)原来被合并成同一句话。"""
        _cite_project(self.tmp, 30, facts=None)
        with io.open(os.path.join(self.tmp, "out", "session_facts.json"), "w",
                     encoding="utf-8", newline=NL) as f:
            f.write(json.dumps({"steps_used": None}))
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertNotIn("不存在", line)
        self.assertIn("没取到", line)

    def test_a_quoted_number_is_not_a_disagreement(self):
        """⛔ `7` 与 `'7'` 判成不符,会产生一条**两边印出来一模一样**的 FAIL,
        读的人只会以为窗口坏了。"""
        _cite_project(self.tmp, 7, facts="7")
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("PASS", line)

    def test_the_machine_number_drives_the_budget_even_when_the_record_is_null(self):
        """⛔ 原来只拿自报的数比上限:dossier 写 null、机器记 200、上限 10,
        报的是 PASS「预算内」。"""
        _cite_project(self.tmp, None, facts=200)
        p = os.path.join(self.tmp, "task_plan.md")
        meta, body = read_md(p)
        meta["budget"] = {"steps_max": 10}
        write_md(p, meta, body)
        out = _run_cite(self.tmp)
        line = [l.strip() for l in out.split(NL) if "预算仪表:" in l][0]
        self.assertIn("WARN", line)
        self.assertIn("200", line)

    def test_a_null_cards_count_does_not_crash_the_whole_check(self):
        """⛔ 与 steps_used 同一形态的 `None > int`:我修了 steps 那一行,把**紧邻
        上一行**的 cards_count 留在原地。崩掉时 stdout 全空 —— 交付对接窗口看起来是
        「什么都没说」而不是「拦住了」。"""
        _cite_project(self.tmp, 5, facts=5)
        d = os.path.join(self.tmp, "dossier.md")
        meta, body = read_md(d)
        meta["dashboard"] = {"cards_count": None, "steps_used": 5}
        write_md(d, meta, body)
        t = os.path.join(self.tmp, "task_plan.md")
        tm, tb = read_md(t)
        tm["budget"] = {"cards_max": 10}
        write_md(t, tm, tb)
        out = _run_cite(self.tmp)
        self.assertNotIn("TypeError", out)
        self.assertIn("预算仪表", out)

    def test_agreeing_with_session_facts_passes(self):
        _cite_project(self.tmp, 255, facts=255)
        self.assertIn("PASS", _steps_line(_run_cite(self.tmp)))

    def test_machine_json_missing_provenance_warns_instead_of_passing(self):
        _cite_project(self.tmp, 255, facts={"steps_used": 255})
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("机器出处本身不可复核", line)

    def test_invalid_machine_capture_time_warns_instead_of_passing(self):
        facts = {"steps_used": 255, "session": "session-a",
                 "session_picked_by": "DSH_SESSION_ID(harness 注入)",
                 "captured_at": "not-a-time"}
        _cite_project(self.tmp, 255, facts=facts,
                      dashboard_extra={"session": facts["session"],
                                       "session_picked_by": facts["session_picked_by"],
                                       "captured_at": facts["captured_at"]})
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("不是带时区的 ISO-8601", line)

    def test_dashboard_using_old_provenance_key_names_warns(self):
        facts = {"steps_used": 255, "session": "session-a",
                 "session_picked_by": "DSH_SESSION_ID(harness 注入)",
                 "captured_at": "2026-09-04T12:00:00.000Z"}
        _cite_project(self.tmp, 255, facts=facts,
                      dashboard_extra={"steps_session": "session-a",
                                       "steps_picked_by": facts["session_picked_by"]})
        line = _steps_line(_run_cite(self.tmp))
        self.assertIn("WARN", line)
        self.assertIn("session / session_picked_by / captured_at", line)

    def test_each_wrong_provenance_field_fails_even_when_steps_match(self):
        facts = {"steps_used": 255, "session": "session-a",
                 "session_picked_by": "DSH_SESSION_ID(harness 注入)",
                 "captured_at": "2026-09-04T12:00:00.000Z"}
        wrong = {"session": "session-b", "session_picked_by": "启发式",
                 "captured_at": "2026-09-04T12:01:00.000Z"}
        for key in ("session", "session_picked_by", "captured_at"):
            with self.subTest(key=key):
                project = tempfile.mkdtemp(prefix="cc-provenance-")
                dashboard = {k: facts[k] for k in
                             ("session", "session_picked_by", "captured_at")}
                dashboard[key] = wrong[key]
                _cite_project(project, 255, facts=facts, dashboard_extra=dashboard)
                line = _steps_line(_run_cite(project))
                self.assertIn("FAIL", line)
                self.assertIn(key + ":", line)



class SessionFactsPicksTheSessionAndSeesSubagents(unittest.TestCase):
    """会话选法、边界判据、子代理归并。

    ⛔ 三个真实缺口(独立复核查出):
      · 启发式「最新一份提到该项目的日志」在五项目实测里**选错 3/5**,而
        harness 每次工具调用都注入了 DSH_SESSION_ID / DSH_SESSION_JSONL。
      · 找不到会话时原来 rc=0 且打「0 条」—— 正是四份 skill 明令禁止的
        「空清单被读成没越界」。
      · 子代理日志完全不在清单里:实测 10 份子代理读了三个项目目录外的文件,
        主会话清单一个都没有。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-sf-")
        self.project = os.path.join(self.tmp, "proj")
        os.makedirs(self.project)
        self.home = os.path.join(self.tmp, "home")
        os.makedirs(os.path.join(self.home, ".dsh", "sessions", "tag"))

    def _call(self, name, args):
        return {"type": "tool/call",
                "data": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}

    def _header(self, sid, parent=None):
        return {"type": "session", "id": sid, "parentSession": parent,
                "origin": "subagent" if parent else None}

    def _log(self, sid, events):
        d = os.path.join(self.home, ".dsh", "sessions", "tag", sid)
        os.makedirs(d, exist_ok=True)
        raw = os.path.join(self.tmp, sid + ".jsonl")
        with io.open(raw, "w", encoding="utf-8", newline="") as f:
            for e in events:
                f.write(json.dumps(e, ensure_ascii=False) + NL)
        out = os.path.join(d, "session.jsonl.zstd")
        script = ("const fs=require('fs'),z=require('zlib');"
                  "fs.writeFileSync(process.argv[2], z.zstdCompressSync(fs.readFileSync(process.argv[1])));")
        r = subprocess.run(["node", "-e", script, "--", raw, out], capture_output=True)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        return out

    def _run(self, project=None, extra_env=None, extra_args=None):
        out = os.path.join(self.tmp, "facts.json")
        if os.path.exists(out):
            os.remove(out)
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home,
                   PYTHONIOENCODING="utf-8")
        env.pop("DSH_SESSION_ID", None)
        env.pop("DSH_SESSION_JSONL", None)
        env.update(extra_env or {})
        argv = ["node", os.path.join(HERE, "session_facts.mjs"),
                "--project", project or self.project, "--json", out]
        argv += list(extra_args or [])
        r = subprocess.run(argv, capture_output=True, env=env)
        text = (r.stdout + r.stderr).decode("utf-8", "replace")
        data = None
        if os.path.exists(out):
            with io.open(out, encoding="utf-8") as f:
                data = json.loads(f.read())
        return r.returncode, text, data

    def test_no_session_for_this_project_exits_nonzero(self):
        """⛔ 「没找到」不许长得像「没越界」:rc≠0,不打 0 条,json 里是 null。"""
        elsewhere = os.path.join(self.tmp, "别处", "x.md")
        self._log("session-a", [self._call("read", {"file_path": elsewhere})])
        code, text, data = self._run()
        self.assertNotEqual(code, 0, text)
        self.assertIn("未取到会话日志", text)
        self.assertNotIn("0 条", text)
        self.assertIsNone(data["steps_used"])
        self.assertIsNone(data["outside_reads"])

    def test_a_sibling_sharing_a_prefix_is_outside(self):
        """⛔ 裸 startsWith 会把 `proj-2` 当成 `proj` 的内部,越界读取凭空少几条。"""
        inside = os.path.join(self.project, "task_plan.md")
        sibling = os.path.join(self.tmp, "proj-2", "task_plan.md")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("read", {"file_path": sibling})])
        _, text, data = self._run()
        paths = [r["path"] for r in data["outside_reads"]]
        self.assertEqual(paths, [sibling], text)

    def test_subagent_reads_are_merged_by_parent_session(self):
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "聊天记录.md")
        self._log("session-p", [self._header("session-p"),
                                self._call("read", {"file_path": inside})])
        self._log("session-k", [self._header("session-k", parent="session-p"),
                                self._call("read", {"file_path": far})])
        _, text, data = self._run()
        paths = [r["path"] for r in data["outside_reads"]]
        self.assertIn(far, paths, text)
        got = [r for r in data["outside_reads"] if r["path"] == far][0]
        self.assertEqual(got["from"], "子代理")
        self.assertEqual(data["subagent_sessions"], ["session-k"])

    def test_grandchild_subagent_is_merged_recursively(self):
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "孙代理读取.md")
        self._log("session-p", [self._header("session-p"),
                                self._call("read", {"file_path": inside})])
        self._log("session-k", [self._header("session-k", parent="session-p")])
        self._log("session-g", [self._header("session-g", parent="session-k"),
                                {"type": "step/start", "data": {}},
                                self._call("read", {"file_path": far})])
        _, text, data = self._run(extra_env={"DSH_SESSION_ID": "session-p"})
        self.assertIn(far, [r["path"] for r in data["outside_reads"]], text)
        self.assertEqual(data["subagent_sessions"], ["session-k", "session-g"])
        self.assertEqual(data["steps_used"], 0)
        self.assertEqual(data["steps_incl_subagents"], 1)

    def test_parent_session_can_refer_to_directory_sid_when_header_id_differs(self):
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "目录会话链.md")
        self._log("session-p", [self._header("logical-root"),
                                self._call("read", {"file_path": inside})])
        self._log("session-k", [self._header("logical-child", parent="session-p"),
                                self._call("read", {"file_path": far})])
        _, text, data = self._run(extra_env={"DSH_SESSION_ID": "session-p"})
        self.assertIn(far, [r["path"] for r in data["outside_reads"]], text)
        self.assertEqual(data["subagent_sessions"], ["session-k"])

    def test_the_harness_variable_beats_the_heuristic(self):
        """harness 每次调用都注入 DSH_SESSION_ID;启发式只该在没有它时兜底。"""
        inside = os.path.join(self.project, "task_plan.md")
        self._log("session-old", [self._call("read", {"file_path": inside})])
        self._log("session-new", [self._call("read", {"file_path": inside})])
        _, _, guessed = self._run()
        _, text, told = self._run(extra_env={"DSH_SESSION_ID": "session-old"})
        self.assertEqual(told["session"], "session-old", text)
        self.assertIn("DSH_SESSION_ID", told["session_picked_by"])
        self.assertIn("启发式", guessed["session_picked_by"])

    def test_real_output_has_capture_time_and_cite_check_accepts_it_verbatim(self):
        inside = os.path.join(self.project, "task_plan.md")
        self._log("session-a", [self._call("read", {"file_path": inside})])
        _, text, data = self._run(extra_env={"DSH_SESSION_ID": "session-a"})
        captured = data.get("captured_at")
        self.assertRegex(captured or "", r"^\d{4}-\d{2}-\d{2}T.*Z$")
        self.assertIn(captured, text)
        dashboard = {k: data[k] for k in ("session", "session_picked_by", "captured_at")}
        _cite_project(self.project, data["steps_used"], facts=data,
                      dashboard_extra=dashboard)
        self.assertIn("PASS", _steps_line(_run_cite(self.project)))

    def test_a_write_is_labelled_a_write(self):
        """⛔ 把 write 标成「读」会把公理 7 违规说成常规读取。"""
        inside = os.path.join(self.project, "task_plan.md")
        written = os.path.join(self.tmp, "别处", "liveness_probe.py")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("write", {"file_path": written})])
        _, text, data = self._run()
        got = [r for r in data["outside_reads"] if r["path"] == written]
        self.assertEqual(len(got), 1, text)
        self.assertEqual(got[0]["action"], "写")

    def test_read_then_write_of_the_same_path_is_labelled_a_write(self):
        """⛔ 路径去重不能把后发生的写入锁死为首次读取。"""
        inside = os.path.join(self.project, "task_plan.md")
        touched = os.path.join(self.tmp, "别处", "checker.py")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("read", {"file_path": touched}),
                                self._call("edit", {"file_path": touched})])
        _, text, data = self._run()
        got = [r for r in data["outside_reads"] if r["path"] == touched]
        self.assertEqual(len(got), 1, text)
        self.assertEqual(got[0]["action"], "写")



class SessionFactsGroupsAndPicksAmongSeveralLogs(unittest.TestCase):
    """清单分组,以及「选哪一份日志」那四段逻辑各自有测试。

    ⛔ 分组的理由:实测归并子代理后 66 条,真正越界的 15 条(项目目录外的资料
    文件 + 子代理落在 Temp 里的 web_fetch 溢写)混在 36 条别的项目、15 条
    library/templates/scripts 中间。平铺等于没列 —— 这条脚本自己写着
    「满是噪声的清单等于没有这一段」。

    ⛔ 选日志的理由:对抗式独立复核指出原来的 fixture 只有一份会话,于是
    「取最新 / 忽略 --session / 任何绝对路径都算提到本项目 / 只解首帧」四个
    变异体全部存活,而前两个在真日志上会换一个答案。
    """

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-sfg-")
        self.toolkit = os.path.join(self.tmp, "toolkit")
        self.project = os.path.join(self.toolkit, "projects", "本项目")
        os.makedirs(self.project)
        os.makedirs(os.path.join(self.toolkit, "projects", "别的项目", "cards"))
        os.makedirs(os.path.join(self.toolkit, "library", "rules"))
        os.makedirs(os.path.join(self.tmp, "桌面"))
        self.home = os.path.join(self.tmp, "home")
        os.makedirs(os.path.join(self.home, ".dsh", "sessions", "tag"))

    def _call(self, name, args):
        return {"type": "tool/call",
                "data": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}

    def _header(self, sid, parent=None):
        return {"type": "session", "id": sid, "parentSession": parent,
                "origin": "subagent" if parent else None}

    def _raw(self, sid, events):
        raw = os.path.join(self.tmp, sid + ".jsonl")
        with io.open(raw, "w", encoding="utf-8", newline="") as f:
            for e in events:
                f.write(json.dumps(e, ensure_ascii=False) + NL)
        return raw

    def _log(self, sid, events, mtime=None, frames=1):
        """写一份会话日志。frames=2 时压成**两个 zstd 帧**再拼接。"""
        d = os.path.join(self.home, ".dsh", "sessions", "tag", sid)
        os.makedirs(d, exist_ok=True)
        out = os.path.join(d, "session.jsonl.zstd")
        raw = self._raw(sid, events)
        if frames == 1:
            script = ("const fs=require('fs'),z=require('zlib');"
                      "fs.writeFileSync(process.argv[2], z.zstdCompressSync(fs.readFileSync(process.argv[1])));")
        else:
            script = ("const fs=require('fs'),z=require('zlib');"
                      "const L=fs.readFileSync(process.argv[1],'utf8').split(String.fromCharCode(10))"
                      ".filter(Boolean);"
                      "const h=Math.ceil(L.length/2);"
                      "const a=L.slice(0,h).join(String.fromCharCode(10))+String.fromCharCode(10);"
                      "const b=L.slice(h).join(String.fromCharCode(10))+String.fromCharCode(10);"
                      "fs.writeFileSync(process.argv[2], Buffer.concat("
                      "[z.zstdCompressSync(Buffer.from(a)), z.zstdCompressSync(Buffer.from(b))]));")
        r = subprocess.run(["node", "-e", script, "--", raw, out], capture_output=True)
        self.assertEqual(r.returncode, 0, r.stderr.decode("utf-8", "replace"))
        if mtime is not None:
            os.utime(out, (mtime, mtime))
        return out

    def _run(self, extra_args=None):
        out = os.path.join(self.tmp, "facts.json")
        if os.path.exists(out):
            os.remove(out)
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home)
        env.pop("DSH_SESSION_ID", None)
        env.pop("DSH_SESSION_JSONL", None)
        argv = ["node", os.path.join(HERE, "session_facts.mjs"),
                "--project", self.project, "--json", out] + list(extra_args or [])
        r = subprocess.run(argv, capture_output=True, env=env)
        text = (r.stdout + r.stderr).decode("utf-8", "replace")
        data = None
        if os.path.exists(out):
            with io.open(out, encoding="utf-8") as f:
                data = json.loads(f.read())
        return r.returncode, text, data

    def _scope_of(self, data, p):
        hit = [r for r in data["outside_reads"] if r["path"] == p]
        self.assertEqual(len(hit), 1, "清单里没有 %s" % p)
        return hit[0]["scope"]

    def test_a_harness_spill_file_is_not_an_escape(self):
        """⛔ `dsh-spill-*` 下的文件是 harness 把工具输出落盘、子代理再读回来 ——
        读的是它自己的工具输出。归进 outside_toolkit 会把真越界稀释掉:实测
        原来 15 条里有 5 条是这个。哨兵:把它归成 outside_toolkit → 红。
        """
        inside = os.path.join(self.project, "task_plan.md")
        spill = os.path.join(self.tmp, "dsh-spill-abc123", "session-ff01", "beef-web_fetch.txt")
        far = os.path.join(self.tmp, "桌面", "聊天记录.md")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("grep", {"path": spill}),
                                self._call("read", {"file_path": far})])
        code, text, data = self._run()
        self.assertEqual(code, 0, text)
        self.assertEqual(self._scope_of(data, spill), "harness_temp", text)
        self.assertEqual(self._scope_of(data, far), "outside_toolkit", text)
        self.assertEqual(data["scope_counts"]["outside_toolkit"], 1, "溢写文件被算成越界了")
        self.assertEqual(data["scope_counts"]["harness_temp"], 1)
        self.assertNotIn(spill, text, "溢写文件不该占控制台的逐条位")

    def test_the_temp_hint_lands_only_on_dsh_dirs_that_are_not_spill(self):
        """提示只是提示,不改分组。⛔ 判据故意窄:放宽成「凡 dsh- 临时目录都不算
        越界」就给出了洗白路径 —— 把项目目录外的文件抄进 Temp/dsh-foo 再读回来,
        清单上就看不见了。所以这类条目留在越界组,只多一行字。
        哨兵:提示落到项目目录外那条或 spill 那条 → 红。
        """
        inside = os.path.join(self.project, "task_plan.md")
        temp_dir = os.path.join(self.tmp, "dsh-AbCdEf", "moncler2020.md")
        spill = os.path.join(self.tmp, "dsh-spill-abc123", "s1", "beef-web_fetch.txt")
        far = os.path.join(self.tmp, "桌面", "聊天记录.md")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("grep", {"path": temp_dir}),
                                self._call("grep", {"path": spill}),
                                self._call("read", {"file_path": far})])
        code, text, data = self._run()
        self.assertEqual(code, 0, text)
        by_path = {r["path"]: r for r in data["outside_reads"]}
        self.assertIsNotNone(by_path[temp_dir]["hint"], "dsh- 非 spill 那条没带提示")
        self.assertEqual(by_path[temp_dir]["scope"], "outside_toolkit", "提示不该改分组")
        self.assertIsNone(by_path[far]["hint"], "项目目录外那条不该带提示")
        self.assertIsNone(by_path[spill]["hint"], "spill 那条不该带提示")
        hinted = [r for r in data["outside_reads"] if r["hint"] is not None]
        self.assertEqual(len(hinted), 1, "带提示的条数不对:%s" % [r["path"] for r in hinted])

    def test_three_scopes_are_told_apart(self):
        """哨兵:把项目目录外那条归错组 → 红。"""
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "聊天记录.md")
        other = os.path.join(self.toolkit, "projects", "别的项目", "cards", "L01.md")
        internal = os.path.join(self.toolkit, "library", "rules", "discipline.md")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("read", {"file_path": far}),
                                self._call("read", {"file_path": other}),
                                self._call("read", {"file_path": internal})])
        code, text, data = self._run()
        self.assertEqual(code, 0, text)
        self.assertEqual(self._scope_of(data, far), "outside_toolkit", text)
        self.assertEqual(self._scope_of(data, other), "other_project", text)
        self.assertEqual(self._scope_of(data, internal), "toolkit_internal", text)
        self.assertEqual(data["scope_counts"]["outside_toolkit"], 1)
        self.assertEqual(len(data["outside_reads"]), 3, "json 折叠了条目 —— 折叠的只该是显示")

    def test_the_console_lists_the_real_ones_and_folds_the_rest(self):
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "聊天记录.md")
        internal = os.path.join(self.toolkit, "library", "rules", "discipline.md")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("read", {"file_path": far}),
                                self._call("read", {"file_path": internal})])
        _, text, _ = self._run()
        self.assertIn(far, text, "真越界那条没逐条列出来")
        self.assertNotIn(internal, text, "工具包内部那条不该占控制台")
        self.assertIn("另 1 条", text)

    def test_the_console_never_folds_a_toolkit_internal_write(self):
        inside = os.path.join(self.project, "task_plan.md")
        internal_read = os.path.join(self.toolkit, "library", "rules", "discipline.md")
        internal_write = os.path.join(self.toolkit, "scripts", "checker.py")
        self._log("session-a", [self._call("read", {"file_path": inside}),
                                self._call("read", {"file_path": internal_read}),
                                self._call("write", {"file_path": internal_write})])
        _, text, data = self._run()
        self.assertIn(internal_write, text, "工具包内部写入被折叠,公理 7 违规不可见")
        self.assertNotIn(internal_read, text, "纯读取仍应折叠,避免把真写入淹没")
        got = [r for r in data["outside_reads"] if r["path"] == internal_write]
        self.assertEqual(got[0]["scope"], "toolkit_internal")
        self.assertEqual(got[0]["action"], "写")
        self.assertIn("另 1 条", text)

    def test_the_newest_log_wins_when_guessing(self):
        inside = os.path.join(self.project, "task_plan.md")
        self._log("session-old", [self._call("read", {"file_path": inside})], mtime=1_000_000)
        self._log("session-new", [self._call("read", {"file_path": inside})], mtime=2_000_000)
        _, text, data = self._run()
        self.assertEqual(data["session"], "session-new", text)

    def test_the_session_flag_overrides_the_guess(self):
        inside = os.path.join(self.project, "task_plan.md")
        self._log("session-old", [self._call("read", {"file_path": inside})], mtime=1_000_000)
        self._log("session-new", [self._call("read", {"file_path": inside})], mtime=2_000_000)
        _, text, data = self._run(extra_args=["--session", "old"])
        self.assertEqual(data["session"], "session-old", text)
        self.assertIn("--session", data["session_picked_by"])

    def test_only_a_path_inside_the_project_counts_as_mentioning_it(self):
        """⛔ 「提到本项目」的判据是路径落在项目**之内**,不是「有绝对路径」。
        放宽成后者,最新那份不相干的日志就会被选中。"""
        inside = os.path.join(self.project, "task_plan.md")
        far = os.path.join(self.tmp, "桌面", "无关.md")
        self._log("session-real", [self._call("read", {"file_path": inside})], mtime=1_000_000)
        self._log("session-noise", [self._call("read", {"file_path": far})], mtime=2_000_000)
        _, text, data = self._run()
        self.assertEqual(data["session"], "session-real", text)

    def test_main_session_steps_and_the_subagent_total_are_kept_apart(self):
        """⛔ 这是 change 3 的头条,却一直没有测试:把两者合成一个数,`steps_used`
        会在有子代理时悄悄翻倍(实测 255 -> 536),而 cite_check ⑧ 拿它当机器
        真值去核 dossier。(对抗核查 2026-09-04。)"""
        inside = os.path.join(self.project, "task_plan.md")
        main = [self._header("session-p"), self._call("read", {"file_path": inside})]
        main += [{"type": "step/start", "data": {}} for _ in range(4)]
        kid = [self._header("session-k", parent="session-p")]
        kid += [{"type": "step/start", "data": {}} for _ in range(6)]
        self._log("session-p", main)
        self._log("session-k", kid)
        _, text, data = self._run()
        self.assertEqual(data["steps_used"], 4, text)
        self.assertEqual(data["steps_incl_subagents"], 10, text)

    def test_a_log_that_decodes_to_nothing_is_not_zero(self):
        """⛔ 选中了却一条事件都解不出来 = 没取到,不是「这一轮什么都没干」。
        原来这条路打印「步数 0 · 项目外路径 0 条」并 exit 0 —— 正是脚本开头
        明令禁止的那种沉默。"""
        bogus = os.path.join(self.tmp, "bogus.jsonl")
        with io.open(bogus, "w", encoding="utf-8", newline=NL) as f:
            f.write("这不是 json,一行也解析不出来" + NL)
        out = os.path.join(self.tmp, "facts.json")
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home,
                   DSH_SESSION_JSONL=bogus)
        env.pop("DSH_SESSION_ID", None)
        r = subprocess.run(["node", os.path.join(HERE, "session_facts.mjs"),
                            "--project", self.project, "--json", out],
                           capture_output=True, env=env)
        text = (r.stdout + r.stderr).decode("utf-8", "replace")
        self.assertNotEqual(r.returncode, 0, text)
        self.assertIn("未取到会话日志", text)
        self.assertNotIn("0 条", text)
        with io.open(out, encoding="utf-8") as f:
            self.assertIsNone(json.loads(f.read())["steps_used"])

    def test_the_json_directory_is_created_instead_of_throwing(self):
        """⛔ 目录不存在时 writeFileSync 抛 ENOENT,node rc=1,而**已经算好的
        结果一行都没打印出来** —— 异常发生在打印之前。"""
        inside = os.path.join(self.project, "task_plan.md")
        self._log("session-a", [self._call("read", {"file_path": inside})])
        deep = os.path.join(self.tmp, "还没有的目录", "再一层", "facts.json")
        env = dict(os.environ, USERPROFILE=self.home, HOME=self.home)
        env.pop("DSH_SESSION_ID", None)
        env.pop("DSH_SESSION_JSONL", None)
        r = subprocess.run(["node", os.path.join(HERE, "session_facts.mjs"),
                            "--project", self.project, "--json", deep],
                           capture_output=True, env=env)
        text = (r.stdout + r.stderr).decode("utf-8", "replace")
        self.assertEqual(r.returncode, 0, text)
        self.assertIn("会话 session-a", text)
        self.assertTrue(os.path.exists(deep), "没建目录:%s" % text)

    def test_every_frame_of_a_multi_frame_log_is_read(self):
        """⛔ 只解首帧会少数步数,而少了的那部分和"本来就少"长得一样。"""
        inside = os.path.join(self.project, "task_plan.md")
        events = [self._call("read", {"file_path": inside})]
        events += [{"type": "step/start", "data": {}} for _ in range(9)]
        self._log("session-a", events, frames=2)
        _, text, data = self._run()
        self.assertEqual(data["steps_used"], 9, text)



class VersionIsReadBeforeItIsCompared(unittest.TestCase):
    """⛔ pipeline_status 是过窗口前第一个跑的东西,它一崩就没有任何一项检查跑过
    —— 而原来那句是裸 `int(m.get("version") or 1)`。(对抗核查 2026-09-04。)"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="pls-ver2-")

    def test_a_non_integer_version_is_reported_not_raised(self):
        p = os.path.join(self.tmp, "task_plan.md")
        with io.open(p, "w", encoding="utf-8", newline=NL) as f:
            f.write("---" + NL + "title: 计划" + NL + "version: 'v2'" + NL
                    + "pipeline_status: planning" + NL + "approval:" + NL
                    + "  status: awaiting" + NL + "---" + NL + NL + "正文。" + NL)
        code, out, err = _run([self.tmp])
        blob = (out + err).decode("utf-8", "replace")
        self.assertNotIn("Traceback", blob)
        self.assertIn("version 不是整数", blob)


if __name__ == "__main__":
    unittest.main(verbosity=2)
